from __future__ import annotations

import pytest
from pptx import Presentation

from src.layout.renderers.cover import CoverRenderer
from src.models.slides import CoverSlide
from src.models.theme import ThemeModel


@pytest.fixture
def slide(blank_presentation):
    layout = blank_presentation.slide_layouts[6]
    return blank_presentation.slides.add_slide(layout)


def test_cover_renders_without_exception(slide, cover_model, minimal_theme) -> None:
    CoverRenderer().render(slide, cover_model, minimal_theme)


def test_cover_minimal_model_renders(slide, minimal_theme) -> None:
    model = CoverSlide(type="cover", title="Minimal")
    CoverRenderer().render(slide, model, minimal_theme)


def test_cover_adds_shapes_to_slide(slide, cover_model, minimal_theme) -> None:
    initial_count = len(slide.shapes)
    CoverRenderer().render(slide, cover_model, minimal_theme)
    assert len(slide.shapes) > initial_count


def test_cover_title_truncation_emits_warning(slide, minimal_theme, caplog) -> None:
    import logging

    model = CoverSlide(type="cover", title="A" * 81)
    with caplog.at_level(logging.WARNING, logger="softgic"):
        CoverRenderer().render(slide, model, minimal_theme)
    assert any("truncated" in r.message.lower() for r in caplog.records)


def test_cover_no_subtitle_renders_cleanly(slide, minimal_theme) -> None:
    model = CoverSlide(type="cover", title="No Subtitle")
    CoverRenderer().render(slide, model, minimal_theme)
