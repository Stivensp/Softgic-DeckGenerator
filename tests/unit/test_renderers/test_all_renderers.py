"""
Unit tests for the 6 remaining layout renderers.
Each renderer is tested for: happy path, minimal model (optional fields absent),
shapes added to slide, and text truncation warning where applicable.
"""
from __future__ import annotations

import logging
from pathlib import Path

import pytest
from PIL import Image
from pptx import Presentation

from src.layout.renderers.closing import ClosingRenderer
from src.layout.renderers.content_two_col import ContentTwoColRenderer
from src.layout.renderers.pricing_table import PricingTableRenderer
from src.layout.renderers.profile_card import ProfileCardRenderer
from src.layout.renderers.section_divider import SectionDividerRenderer
from src.layout.renderers.stat_callout import StatCalloutRenderer
from src.models.slides import (
    BulletColumn,
    ClosingSlide,
    ContentTwoColSlide,
    PricingRow,
    PricingTableSlide,
    ProfileCardSlide,
    SectionDividerSlide,
    StatCalloutSlide,
)


@pytest.fixture
def slide(blank_presentation):
    layout = blank_presentation.slide_layouts[6]
    return blank_presentation.slides.add_slide(layout)


# ── SectionDivider ─────────────────────────────────────────────────────────────

def test_section_divider_renders(slide, section_divider_model, minimal_theme) -> None:
    SectionDividerRenderer().render(slide, section_divider_model, minimal_theme)


def test_section_divider_minimal(slide, minimal_theme) -> None:
    model = SectionDividerSlide(type="section_divider", section_number="01", section_title="Intro")
    SectionDividerRenderer().render(slide, model, minimal_theme)


def test_section_divider_adds_shapes(slide, section_divider_model, minimal_theme) -> None:
    initial = len(slide.shapes)
    SectionDividerRenderer().render(slide, section_divider_model, minimal_theme)
    assert len(slide.shapes) > initial


def test_section_divider_tagline_absent(slide, minimal_theme) -> None:
    model = SectionDividerSlide(
        type="section_divider", section_number="02", section_title="Sin tagline"
    )
    SectionDividerRenderer().render(slide, model, minimal_theme)


def test_section_divider_title_truncation(slide, minimal_theme, caplog) -> None:
    model = SectionDividerSlide(
        type="section_divider", section_number="03", section_title="A" * 81
    )
    with caplog.at_level(logging.WARNING, logger="softgic"):
        SectionDividerRenderer().render(slide, model, minimal_theme)
    assert any("truncated" in r.message.lower() for r in caplog.records)


# ── ContentTwoCol ──────────────────────────────────────────────────────────────

def test_content_two_col_renders(slide, content_two_col_model, minimal_theme) -> None:
    ContentTwoColRenderer().render(slide, content_two_col_model, minimal_theme)


def test_content_two_col_without_column_titles(slide, minimal_theme) -> None:
    model = ContentTwoColSlide(
        type="content_two_col",
        title="Plain",
        left=BulletColumn(bullets=["L1", "L2"]),
        right=BulletColumn(bullets=["R1"]),
    )
    ContentTwoColRenderer().render(slide, model, minimal_theme)


def test_content_two_col_adds_shapes(slide, content_two_col_model, minimal_theme) -> None:
    initial = len(slide.shapes)
    ContentTwoColRenderer().render(slide, content_two_col_model, minimal_theme)
    assert len(slide.shapes) > initial


def test_content_two_col_bullet_truncation(slide, minimal_theme, caplog) -> None:
    model = ContentTwoColSlide(
        type="content_two_col",
        title="T",
        left=BulletColumn(bullets=["X" * 101]),
        right=BulletColumn(bullets=["short"]),
    )
    with caplog.at_level(logging.WARNING, logger="softgic"):
        ContentTwoColRenderer().render(slide, model, minimal_theme)
    assert any("truncated" in r.message.lower() for r in caplog.records)


# ── PricingTable ───────────────────────────────────────────────────────────────

def test_pricing_table_renders(slide, pricing_table_model, minimal_theme) -> None:
    PricingTableRenderer().render(slide, pricing_table_model, minimal_theme)


def test_pricing_table_multiple_rows(slide, minimal_theme) -> None:
    model = PricingTableSlide(
        type="pricing_table",
        title="Multi",
        currency="COP",
        rows=[
            PricingRow(description=f"Item {i}", quantity=i, unit_price=100.0, total=i * 100.0)
            for i in range(1, 6)
        ],
        totals=1500.0,
    )
    PricingTableRenderer().render(slide, model, minimal_theme)


def test_pricing_table_with_notes(slide, minimal_theme) -> None:
    model = PricingTableSlide(
        type="pricing_table",
        title="Con notas",
        currency="USD",
        rows=[PricingRow(description="X", quantity=1, unit_price=50.0, total=50.0)],
        totals=50.0,
        notes="Sin IVA.",
    )
    PricingTableRenderer().render(slide, model, minimal_theme)


def test_pricing_table_without_notes(slide, minimal_theme) -> None:
    model = PricingTableSlide(
        type="pricing_table",
        title="Sin notas",
        currency="USD",
        rows=[PricingRow(description="Y", quantity=1, unit_price=10.0, total=10.0)],
        totals=10.0,
    )
    PricingTableRenderer().render(slide, model, minimal_theme)


def test_pricing_table_adds_table_shape(slide, pricing_table_model, minimal_theme) -> None:
    initial = len(slide.shapes)
    PricingTableRenderer().render(slide, pricing_table_model, minimal_theme)
    assert len(slide.shapes) > initial


# ── ProfileCard ────────────────────────────────────────────────────────────────

def test_profile_card_renders(slide, profile_card_model, minimal_theme) -> None:
    ProfileCardRenderer().render(slide, profile_card_model, minimal_theme)


def test_profile_card_no_photo_uses_placeholder(slide, minimal_theme) -> None:
    model = ProfileCardSlide(
        type="profile_card",
        name="Sin Foto",
        role="Dev",
        years_experience=3,
        seniority="Mid",
        skills=["Python"],
    )
    ProfileCardRenderer().render(slide, model, minimal_theme)


def test_profile_card_with_real_photo(tmp_path, slide, minimal_theme) -> None:
    photo = tmp_path / "foto.jpg"
    img = Image.new("RGB", (400, 400), color=(100, 150, 200))
    img.save(str(photo))

    model = ProfileCardSlide(
        type="profile_card",
        name="Con Foto",
        role="Dev",
        years_experience=5,
        seniority="Senior",
        skills=["Python", "AWS"],
        photo=str(photo),
    )
    ProfileCardRenderer().render(slide, model, minimal_theme)


def test_profile_card_no_highlights(slide, minimal_theme) -> None:
    model = ProfileCardSlide(
        type="profile_card",
        name="Sin Highlights",
        role="QA",
        years_experience=2,
        seniority="Junior",
        skills=["Selenium"],
    )
    ProfileCardRenderer().render(slide, model, minimal_theme)


def test_profile_card_max_skills_capped(slide, minimal_theme) -> None:
    model = ProfileCardSlide(
        type="profile_card",
        name="Muchas Skills",
        role="Dev",
        years_experience=10,
        seniority="Senior",
        skills=["s1", "s2", "s3", "s4", "s5", "s6", "s7", "s8", "s9", "s10"],
    )
    ProfileCardRenderer().render(slide, model, minimal_theme)


def test_profile_card_adds_shapes(slide, profile_card_model, minimal_theme) -> None:
    initial = len(slide.shapes)
    ProfileCardRenderer().render(slide, profile_card_model, minimal_theme)
    assert len(slide.shapes) > initial


# ── StatCallout ────────────────────────────────────────────────────────────────

def test_stat_callout_renders(slide, stat_callout_model, minimal_theme) -> None:
    StatCalloutRenderer().render(slide, stat_callout_model, minimal_theme)


def test_stat_callout_minimal(slide, minimal_theme) -> None:
    model = StatCalloutSlide(type="stat_callout", big_number="+300", label="Proyectos")
    StatCalloutRenderer().render(slide, model, minimal_theme)


def test_stat_callout_with_context_and_source(slide, minimal_theme) -> None:
    model = StatCalloutSlide(
        type="stat_callout",
        big_number="99.9%",
        label="Uptime garantizado",
        context="Medido en producción durante 2025",
        source="Área de Infraestructura",
    )
    StatCalloutRenderer().render(slide, model, minimal_theme)


def test_stat_callout_adds_shapes(slide, stat_callout_model, minimal_theme) -> None:
    initial = len(slide.shapes)
    StatCalloutRenderer().render(slide, stat_callout_model, minimal_theme)
    assert len(slide.shapes) > initial


def test_stat_callout_label_truncation(slide, minimal_theme, caplog) -> None:
    model = StatCalloutSlide(type="stat_callout", big_number="42", label="L" * 81)
    with caplog.at_level(logging.WARNING, logger="softgic"):
        StatCalloutRenderer().render(slide, model, minimal_theme)
    assert any("truncated" in r.message.lower() for r in caplog.records)


# ── Closing ────────────────────────────────────────────────────────────────────

def test_closing_renders(slide, closing_model, minimal_theme) -> None:
    ClosingRenderer().render(slide, closing_model, minimal_theme)


def test_closing_minimal(slide, minimal_theme) -> None:
    model = ClosingSlide(
        type="closing",
        headline="Hablemos",
        contact_name="Juan",
        contact_email="juan@softgic.com",
    )
    ClosingRenderer().render(slide, model, minimal_theme)


def test_closing_with_cta(slide, minimal_theme) -> None:
    model = ClosingSlide(
        type="closing",
        headline="¿Listo?",
        contact_name="Ana",
        contact_email="ana@softgic.com",
        contact_phone="+57 300 000 0000",
        cta="Agenda tu demo ahora",
    )
    ClosingRenderer().render(slide, model, minimal_theme)


def test_closing_without_cta_or_phone(slide, minimal_theme) -> None:
    model = ClosingSlide(
        type="closing",
        headline="Gracias",
        contact_name="Carlos",
        contact_email="carlos@softgic.com",
    )
    ClosingRenderer().render(slide, model, minimal_theme)


def test_closing_headline_truncation(slide, minimal_theme, caplog) -> None:
    model = ClosingSlide(
        type="closing",
        headline="H" * 81,
        contact_name="N",
        contact_email="n@softgic.com",
    )
    with caplog.at_level(logging.WARNING, logger="softgic"):
        ClosingRenderer().render(slide, model, minimal_theme)
    assert any("truncated" in r.message.lower() for r in caplog.records)


def test_closing_adds_shapes(slide, closing_model, minimal_theme) -> None:
    initial = len(slide.shapes)
    ClosingRenderer().render(slide, closing_model, minimal_theme)
    assert len(slide.shapes) > initial
