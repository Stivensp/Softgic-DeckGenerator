from __future__ import annotations

import lxml.etree as etree
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

from src.ooxml_engine.parser.text_parser import parse_text_frame


def test_run_font_properties(blank_slide):
    tb = blank_slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
    run = tb.text_frame.paragraphs[0].add_run()
    run.text = "Hola mundo"
    run.font.bold = True
    run.font.italic = True
    run.font.underline = True
    run.font.size = Pt(24)
    run.font.color.rgb = RGBColor(0x00, 0xAE, 0xEF)
    run.font.name = "Calibri"

    paras = parse_text_frame(tb)

    assert len(paras) == 1
    font = paras[0].runs[0].font
    assert paras[0].runs[0].text == "Hola mundo"
    assert font.bold is True
    assert font.italic is True
    assert font.underline is True
    assert font.size_pt == 24.0
    assert font.color_hex == "00AEEF"
    assert font.family == "Calibri"


def test_multiple_paragraphs_and_alignment(blank_slide):
    tb = blank_slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(2))
    tf = tb.text_frame
    tf.paragraphs[0].add_run().text = "Primera linea"
    tf.paragraphs[0].alignment = PP_ALIGN.CENTER
    p2 = tf.add_paragraph()
    p2.add_run().text = "Segunda linea"
    p2.alignment = PP_ALIGN.RIGHT

    paras = parse_text_frame(tb)

    assert len(paras) == 2
    assert paras[0].alignment == "center"
    assert paras[0].runs[0].text == "Primera linea"
    assert paras[1].alignment == "right"
    assert paras[1].runs[0].text == "Segunda linea"


def test_bullet_char_detected(blank_slide):
    tb = blank_slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
    p = tb.text_frame.paragraphs[0]
    p.add_run().text = "Item con vineta"
    ppr = p._p.get_or_add_pPr()
    bu_char = etree.SubElement(ppr, qn("a:buChar"))
    bu_char.set("char", "-")

    paras = parse_text_frame(tb)

    assert paras[0].bullet_char == "-"


def test_empty_text_frame_returns_no_paragraphs(blank_slide):
    tb = blank_slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))

    paras = parse_text_frame(tb)

    assert paras == []


def test_multirun_paragraph_preserves_each_run(blank_slide):
    tb = blank_slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
    p = tb.text_frame.paragraphs[0]
    r1 = p.add_run()
    r1.text = "Normal "
    r2 = p.add_run()
    r2.text = "Negrita"
    r2.font.bold = True

    paras = parse_text_frame(tb)

    assert len(paras[0].runs) == 2
    assert paras[0].runs[0].text == "Normal "
    assert paras[0].runs[0].font.bold is False
    assert paras[0].runs[1].text == "Negrita"
    assert paras[0].runs[1].font.bold is True
