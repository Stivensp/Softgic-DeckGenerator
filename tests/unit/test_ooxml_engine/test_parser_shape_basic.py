from __future__ import annotations

from PIL import Image
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches

from src.ooxml_engine.model.elements import ElementKind
from src.ooxml_engine.model.style import FillKind
from src.ooxml_engine.parser.shape_parser import parse_shape


def test_solid_rectangle_no_text(blank_slide):
    rect = blank_slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1), Inches(2), Inches(3), Inches(1))
    rect.fill.solid()
    rect.fill.fore_color.rgb = __import__("pptx.dml.color", fromlist=["RGBColor"]).RGBColor(0xFF, 0x00, 0x00)
    rect.line.fill.background()

    el = parse_shape(rect)

    assert el.kind == ElementKind.SHAPE
    assert el.fill.kind == FillKind.SOLID
    assert el.fill.color_hex == "FF0000"
    assert el.text is None
    assert el.transform.left_emu == Inches(1)
    assert el.transform.top_emu == Inches(2)
    assert el.transform.width_emu == Inches(3)
    assert el.transform.height_emu == Inches(1)


def test_rectangle_with_fill_and_text_keeps_both(blank_slide):
    rect = blank_slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1), Inches(1), Inches(2), Inches(0.5))
    rect.fill.solid()
    rect.fill.fore_color.rgb = __import__("pptx.dml.color", fromlist=["RGBColor"]).RGBColor(0x00, 0xAE, 0xEF)
    rect.text_frame.text = "Agenda tu demo"

    el = parse_shape(rect)

    assert el.kind == ElementKind.SHAPE
    assert el.fill.color_hex == "00AEEF"
    assert el.text is not None
    assert el.text.paragraphs[0].runs[0].text == "Agenda tu demo"


def test_rectangle_no_fill_with_text_becomes_text_element(blank_slide):
    rect = blank_slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1), Inches(1), Inches(2), Inches(0.5))
    rect.fill.background()
    rect.text_frame.text = "Solo texto, sin relleno"

    el = parse_shape(rect)

    assert el.kind == ElementKind.TEXT
    assert el.paragraphs[0].runs[0].text == "Solo texto, sin relleno"


def test_plain_textbox(blank_slide):
    tb = blank_slide.shapes.add_textbox(Inches(0.5), Inches(0.5), Inches(4), Inches(1))
    tb.text_frame.text = "Titulo del slide"

    el = parse_shape(tb)

    assert el.kind == ElementKind.TEXT
    assert el.paragraphs[0].runs[0].text == "Titulo del slide"


def test_rotation_captured(blank_slide):
    rect = blank_slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1), Inches(1), Inches(2), Inches(1))
    rect.rotation = 45.0

    el = parse_shape(rect)

    assert el.transform.rotation_deg == 45.0


def test_picture_element(blank_slide, tmp_path):
    photo = tmp_path / "foto.png"
    Image.new("RGB", (200, 200), color=(10, 20, 30)).save(str(photo))
    pic = blank_slide.shapes.add_picture(str(photo), Inches(1), Inches(1), Inches(2), Inches(2))

    el = parse_shape(pic)

    assert el.kind == ElementKind.IMAGE
    assert el.image_format == "png"
    assert len(el.blob) > 0


def test_empty_shape_has_no_paragraphs(blank_slide):
    rect = blank_slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1), Inches(1), Inches(2), Inches(1))
    rect.fill.solid()
    rect.fill.fore_color.rgb = __import__("pptx.dml.color", fromlist=["RGBColor"]).RGBColor(0x00, 0x00, 0x00)

    el = parse_shape(rect)

    assert el.kind == ElementKind.SHAPE
    assert el.text is None
