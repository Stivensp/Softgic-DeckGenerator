from __future__ import annotations

from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
from pptx.util import Inches

from src.ooxml_engine.parser.xml_readers import (
    get_prst_geometry,
    read_autofit_font_scale,
    read_fill_hex,
    read_text_margins_emu,
)


def test_get_prst_geometry_known_shape(blank_slide):
    rect = blank_slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1), Inches(1), Inches(1), Inches(1))
    assert get_prst_geometry(rect) == "rect"


def test_get_prst_geometry_unrecognized_prst_does_not_raise(blank_slide):
    """Replica el bug ya corregido esta sesion en apply_visual_config.py:
    shape.auto_shape_type lanza ValueError (no devuelve None) cuando el
    prst no esta en el enum de python-pptx. get_prst_geometry lee el XML
    crudo directamente y nunca pasa por esa propiedad, asi que no deberia
    verse afectado — esta prueba lo confirma explicitamente.
    """
    rect = blank_slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1), Inches(1), Inches(1), Inches(1))
    prst_geom = rect._element.spPr.find(qn("a:prstGeom"))
    prst_geom.set("prst", "totallyBogusShapeName")

    assert get_prst_geometry(rect) == "totallyBogusShapeName"


def test_get_prst_geometry_textbox_defaults_to_rect(blank_slide):
    """Un textbox de python-pptx es, a nivel OOXML, un auto-shape con
    txBox="1" — SI tiene prstGeom (prst="rect" por defecto), no None. La
    diferencia con un AUTO_SHAPE normal es el atributo txBox, no la
    ausencia de geometria."""
    tb = blank_slide.shapes.add_textbox(Inches(1), Inches(1), Inches(1), Inches(1))
    assert get_prst_geometry(tb) == "rect"


def test_read_fill_hex_solid(blank_slide):
    rect = blank_slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1), Inches(1), Inches(1), Inches(1))
    rect.fill.solid()
    rect.fill.fore_color.rgb = RGBColor(0x12, 0x34, 0x56)

    assert read_fill_hex(rect) == "123456"


def test_read_fill_hex_no_fill_returns_none(blank_slide):
    rect = blank_slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1), Inches(1), Inches(1), Inches(1))
    rect.fill.background()

    assert read_fill_hex(rect) is None


def test_read_autofit_font_scale_default_is_one(blank_slide):
    tb = blank_slide.shapes.add_textbox(Inches(1), Inches(1), Inches(1), Inches(1))
    tb.text_frame.text = "sin autofit"

    assert read_autofit_font_scale(tb) == 1.0


def test_read_autofit_font_scale_reads_normautofit(blank_slide):
    tb = blank_slide.shapes.add_textbox(Inches(1), Inches(1), Inches(1), Inches(1))
    tb.text_frame.text = "con autofit"
    body_pr = tb.text_frame._txBody.find(qn("a:bodyPr"))
    norm_autofit = body_pr.makeelement(qn("a:normAutofit"), {"fontScale": "62000"})
    body_pr.append(norm_autofit)

    assert read_autofit_font_scale(tb) == 0.62


def test_read_text_margins_emu_returns_pptx_defaults(blank_slide):
    tb = blank_slide.shapes.add_textbox(Inches(1), Inches(1), Inches(1), Inches(1))

    left, top, right, bottom = read_text_margins_emu(tb)

    assert left == Inches(0.1)
    assert right == Inches(0.1)
    assert top == Inches(0.05)
    assert bottom == Inches(0.05)
