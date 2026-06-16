from __future__ import annotations

import logging

import pytest
from pptx import Presentation

from src.exceptions import RenderError, UnregisteredLayoutError
from src.layout.base_layout import BaseLayoutRenderer
from src.layout.registry import LayoutRegistry
from src.models.theme import ThemeModel
from src.renderer.deck_renderer import DeckRenderer


class _GoodRenderer(BaseLayoutRenderer):
    def render(self, slide: object, model: object, theme: ThemeModel) -> None:
        pass


class _BrokenRenderer(BaseLayoutRenderer):
    def render(self, slide: object, model: object, theme: ThemeModel) -> None:
        raise RuntimeError("renderer exploded")


@pytest.fixture
def registry() -> LayoutRegistry:
    reg = LayoutRegistry()
    reg.register("cover", _GoodRenderer)
    return reg


def test_renderer_returns_presentation(full_deck, minimal_theme, registry) -> None:
    registry.register("section_divider", _GoodRenderer)
    registry.register("content_two_col", _GoodRenderer)
    registry.register("pricing_table", _GoodRenderer)
    registry.register("profile_card", _GoodRenderer)
    registry.register("stat_callout", _GoodRenderer)
    registry.register("closing", _GoodRenderer)

    prs = Presentation()
    result = DeckRenderer().render(full_deck, minimal_theme, registry, prs)
    assert result is prs
    assert len(result.slides) == 7


def test_renderer_raises_render_error_on_failure(minimal_theme, caplog) -> None:
    from src.models.deck import DeckModel
    from src.models.slides import CoverSlide

    reg = LayoutRegistry()
    reg.register("cover", _BrokenRenderer)
    deck = DeckModel(slides=[CoverSlide(type="cover", title="T")])
    prs = Presentation()

    with pytest.raises(RenderError, match="renderer exploded"):
        DeckRenderer().render(deck, minimal_theme, reg, prs)


def test_renderer_layout_index_fallback_logs_warning(minimal_theme, caplog) -> None:
    from src.models.deck import DeckModel
    from src.models.slides import CoverSlide
    from src.models.theme import SlideLayoutIndex

    # Force a layout_indices.cover value beyond what a blank Presentation has
    import copy
    from pydantic import ConfigDict

    # Build a theme with a very high layout index
    data = minimal_theme.model_dump()
    data["layout_indices"]["cover"] = 999
    from src.models.theme import ThemeModel
    high_theme = ThemeModel.model_validate(data)

    reg = LayoutRegistry()
    reg.register("cover", _GoodRenderer)
    deck = DeckModel(slides=[CoverSlide(type="cover", title="T")])
    prs = Presentation()

    with caplog.at_level(logging.WARNING, logger="softgic"):
        DeckRenderer().render(deck, high_theme, reg, prs)

    assert any("exceeds" in r.message.lower() for r in caplog.records)
