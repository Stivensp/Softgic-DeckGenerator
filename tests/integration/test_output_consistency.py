from __future__ import annotations

import pytest
from pptx import Presentation

from src.exporter.pptx_exporter import PptxExporter
from src.layout.registry import LayoutRegistry
from src.layout.renderers.closing import ClosingRenderer
from src.layout.renderers.content_two_col import ContentTwoColRenderer
from src.layout.renderers.cover import CoverRenderer
from src.layout.renderers.pricing_table import PricingTableRenderer
from src.layout.renderers.profile_card import ProfileCardRenderer
from src.layout.renderers.section_divider import SectionDividerRenderer
from src.layout.renderers.stat_callout import StatCalloutRenderer
from src.parser.deck_parser import DeckParser
from src.renderer.deck_renderer import DeckRenderer


@pytest.fixture
def registry() -> LayoutRegistry:
    reg = LayoutRegistry()
    reg.register("cover", CoverRenderer)
    reg.register("section_divider", SectionDividerRenderer)
    reg.register("content_two_col", ContentTwoColRenderer)
    reg.register("pricing_table", PricingTableRenderer)
    reg.register("profile_card", ProfileCardRenderer)
    reg.register("stat_callout", StatCalloutRenderer)
    reg.register("closing", ClosingRenderer)
    return reg


def _build_pptx(raw: dict, theme, registry) -> bytes:
    """Run the pipeline and return the output file bytes."""
    import io

    deck = DeckParser().parse(raw)
    prs = Presentation()
    prs = DeckRenderer().render(deck, theme, registry, prs)
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def test_same_input_same_output_size(full_raw_deck, minimal_theme, registry) -> None:
    """N04: same input → same output (byte size must be identical across two runs)."""
    out1 = _build_pptx(full_raw_deck, minimal_theme, registry)
    out2 = _build_pptx(full_raw_deck, minimal_theme, registry)
    assert len(out1) == len(out2)


def test_different_slide_count_different_output(minimal_theme, registry) -> None:
    short = {"slides": [{"type": "cover", "title": "Short"}]}
    longer = {
        "slides": [
            {"type": "cover", "title": "Longer"},
            {"type": "stat_callout", "big_number": "99%", "label": "L"},
        ]
    }
    out_short = _build_pptx(short, minimal_theme, registry)
    out_longer = _build_pptx(longer, minimal_theme, registry)
    assert len(out_short) != len(out_longer)


def test_output_is_valid_pptx(tmp_path, full_raw_deck, minimal_theme, registry) -> None:
    output = tmp_path / "verify.pptx"
    deck = DeckParser().parse(full_raw_deck)
    prs = Presentation()
    prs = DeckRenderer().render(deck, minimal_theme, registry, prs)
    PptxExporter().export(prs, output)

    # python-pptx can open its own output without errors
    reloaded = Presentation(str(output))
    assert len(reloaded.slides) == 7
