from __future__ import annotations

from pathlib import Path

import pytest
from pptx import Presentation

from src.exporter.pptx_exporter import PptxExporter
from src.layout.registry import LayoutRegistry
from src.layout.renderers.closing import ClosingRenderer
from src.layout.renderers.content_two_col import ContentTwoColRenderer
from src.layout.renderers.cover import CoverRenderer
from src.layout.renderers.pricing_table import PricingTableRenderer
from src.layout.renderers.profile_card import ProfileCardRenderer
from src.layout.renderers.section_divider import SectionDividerRenderer
from src.layout.renderers.stat_callout import StatCalloutRenderer
from src.parser.deck_parser import DeckParser
from src.renderer.deck_renderer import DeckRenderer
from src.validator.deck_validator import DeckValidator


@pytest.fixture
def registry() -> LayoutRegistry:
    reg = LayoutRegistry()
    reg.register("cover", CoverRenderer)
    reg.register("section_divider", SectionDividerRenderer)
    reg.register("content_two_col", ContentTwoColRenderer)
    reg.register("pricing_table", PricingTableRenderer)
    reg.register("profile_card", ProfileCardRenderer)
    reg.register("stat_callout", StatCalloutRenderer)
    reg.register("closing", ClosingRenderer)
    return reg


def test_full_pipeline_produces_pptx(tmp_path, full_raw_deck, minimal_theme, registry) -> None:
    output = tmp_path / "output.pptx"

    DeckValidator().validate(full_raw_deck)
    deck = DeckParser().parse(full_raw_deck)
    prs = Presentation()
    prs = DeckRenderer().render(deck, minimal_theme, registry, prs)
    PptxExporter().export(prs, output)

    assert output.exists()
    assert output.stat().st_size > 0


def test_full_pipeline_correct_slide_count(tmp_path, full_raw_deck, minimal_theme, registry) -> None:
    deck = DeckParser().parse(full_raw_deck)
    prs = Presentation()
    prs = DeckRenderer().render(deck, minimal_theme, registry, prs)

    assert len(prs.slides) == len(full_raw_deck["slides"])


def test_full_pipeline_all_seven_types(tmp_path, full_raw_deck, minimal_theme, registry) -> None:
    DeckValidator().validate(full_raw_deck)
    deck = DeckParser().parse(full_raw_deck)
    prs = Presentation()
    prs = DeckRenderer().render(deck, minimal_theme, registry, prs)

    assert len(prs.slides) == 7


def test_validation_blocks_render_on_invalid_input(minimal_theme, registry) -> None:
    from src.exceptions import SchemaValidationError

    bad_input = {"slides": [{"type": "cover"}]}  # missing title
    with pytest.raises(SchemaValidationError):
        DeckValidator().validate(bad_input)


def test_export_creates_file_in_new_directory(tmp_path, full_raw_deck, minimal_theme, registry) -> None:
    output = tmp_path / "subdir" / "nested" / "deck.pptx"

    deck = DeckParser().parse(full_raw_deck)
    prs = Presentation()
    prs = DeckRenderer().render(deck, minimal_theme, registry, prs)
    PptxExporter().export(prs, output)

    assert output.exists()
