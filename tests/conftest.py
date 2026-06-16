from __future__ import annotations

import pytest
from pptx import Presentation

from src.models.deck import DeckModel, DeckMetadata
from src.models.slides import (
    BulletColumn,
    ClosingSlide,
    ContentTwoColSlide,
    CoverSlide,
    PricingRow,
    PricingTableSlide,
    ProfileCardSlide,
    SectionDividerSlide,
    StatCalloutSlide,
)
from src.models.theme import (
    AssetPaths,
    ColorPalette,
    FontConfig,
    Limits,
    SlideLayoutIndex,
    ThemeModel,
)


@pytest.fixture
def minimal_theme() -> ThemeModel:
    return ThemeModel(
        version="0.0.1-test",
        colors=ColorPalette(
            primary="#0A1628",
            secondary="#1E3A5F",
            accent="#00AEEF",
            background="#FFFFFF",
            text_dark="#0A1628",
            text_light="#FFFFFF",
            text_muted="#6B7280",
            divider="#E5E7EB",
        ),
        fonts=FontConfig(
            family="Calibri",
            size_heading=36,
            size_subheading=24,
            size_body=14,
            size_caption=10,
            size_stat=72,
            bold_headings=True,
        ),
        assets=AssetPaths(
            template="assets/template.pptx",
            logo_white="assets/logo_white.png",
            logo_dark="assets/logo_dark.png",
            placeholder_profile="assets/placeholder_profile.png",
        ),
        limits=Limits(
            title_max_chars=80,
            subtitle_max_chars=120,
            bullet_max_chars=100,
            bullets_per_column_max=6,
            skills_max=8,
            pricing_rows_max=10,
            profile_image_min_width=200,
            profile_image_min_height=200,
        ),
        layout_indices=SlideLayoutIndex(
            cover=6,
            section_divider=6,
            content_two_col=6,
            pricing_table=6,
            profile_card=6,
            stat_callout=6,
            closing=6,
        ),
    )


@pytest.fixture
def blank_presentation() -> Presentation:
    return Presentation()


@pytest.fixture
def blank_slide(blank_presentation: Presentation):
    layout = blank_presentation.slide_layouts[6]  # blank layout
    return blank_presentation.slides.add_slide(layout)


@pytest.fixture
def cover_model() -> CoverSlide:
    return CoverSlide(
        type="cover",
        title="Test Title",
        subtitle="Test Subtitle",
        client="ACME Corp",
        date="Junio 2026",
        author="Test Author",
    )


@pytest.fixture
def section_divider_model() -> SectionDividerSlide:
    return SectionDividerSlide(
        type="section_divider",
        section_number="01",
        section_title="Introducción",
        tagline="Todo empieza con el contexto.",
    )


@pytest.fixture
def content_two_col_model() -> ContentTwoColSlide:
    return ContentTwoColSlide(
        type="content_two_col",
        title="Comparativa",
        left=BulletColumn(title="Antes", bullets=["Bullet 1", "Bullet 2"]),
        right=BulletColumn(title="Después", bullets=["Result 1", "Result 2"]),
    )


@pytest.fixture
def pricing_table_model() -> PricingTableSlide:
    return PricingTableSlide(
        type="pricing_table",
        title="Inversión",
        currency="USD",
        rows=[PricingRow(description="Análisis", quantity=1, unit_price=1000.0, total=1000.0)],
        totals=1000.0,
        notes="Sin IVA",
    )


@pytest.fixture
def profile_card_model() -> ProfileCardSlide:
    return ProfileCardSlide(
        type="profile_card",
        name="Ana García",
        role="Senior Developer",
        years_experience=7,
        seniority="Senior",
        skills=["Python", "AWS", "Docker"],
        highlights=["Led migration to microservices"],
    )


@pytest.fixture
def stat_callout_model() -> StatCalloutSlide:
    return StatCalloutSlide(
        type="stat_callout",
        big_number="98%",
        label="Satisfacción del cliente",
        context="Basado en encuestas 2025",
        source="Customer Success",
    )


@pytest.fixture
def closing_model() -> ClosingSlide:
    return ClosingSlide(
        type="closing",
        headline="¿Listo para empezar?",
        contact_name="Carlos G.",
        contact_email="carlos@softgic.com",
        contact_phone="+57 300 123 4567",
        cta="Agenda tu demo hoy",
    )


@pytest.fixture
def full_deck(
    cover_model,
    section_divider_model,
    content_two_col_model,
    pricing_table_model,
    profile_card_model,
    stat_callout_model,
    closing_model,
) -> DeckModel:
    return DeckModel(
        slides=[
            cover_model,
            section_divider_model,
            content_two_col_model,
            pricing_table_model,
            profile_card_model,
            stat_callout_model,
            closing_model,
        ]
    )


@pytest.fixture
def minimal_raw_deck() -> dict:
    return {
        "slides": [
            {"type": "cover", "title": "Deck mínimo"},
        ]
    }


@pytest.fixture
def full_raw_deck() -> dict:
    return {
        "slides": [
            {"type": "cover", "title": "Full Deck", "subtitle": "Sub"},
            {
                "type": "section_divider",
                "section_number": "01",
                "section_title": "Intro",
            },
            {
                "type": "content_two_col",
                "title": "Two Cols",
                "left": {"title": "L", "bullets": ["L1"]},
                "right": {"title": "R", "bullets": ["R1"]},
            },
            {
                "type": "pricing_table",
                "title": "Precios",
                "currency": "USD",
                "rows": [
                    {"description": "Item", "quantity": 1, "unit_price": 100.0, "total": 100.0}
                ],
                "totals": 100.0,
            },
            {
                "type": "profile_card",
                "name": "Ana",
                "role": "Dev",
                "years_experience": 5,
                "seniority": "Mid",
                "skills": ["Python"],
            },
            {"type": "stat_callout", "big_number": "99%", "label": "Uptime"},
            {
                "type": "closing",
                "headline": "Hablemos",
                "contact_name": "Juan",
                "contact_email": "juan@softgic.com",
            },
        ]
    }
