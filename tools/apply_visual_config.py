"""
Lee config/config_visual.pptx (editado en PowerPoint) y actualiza
config/generated/layout.yaml con las nuevas posiciones, tamanoss de fuente y colores.

Que lee por tipo de elemento:
  - Rectangulos / ovaIos de color: posicion + color de relleno + kind (box|oval)
  - Cajas de texto:                posicion + tamano de fuente + color del texto

Logica de ocultamiento (el editor visual es la fuente de verdad — lo que no
esta en el PPTX, no aparece en el deck):
  - Elemento en _D del renderer pero NO en el PPTX → {hidden: true} en layout.yaml
  - Elemento extra (no en _D) presente en PPTX → guardado con sus props
  - Elemento extra en old_section pero NO en PPTX → {hidden: true}
  - Elemento previamente {hidden: true} que SI aparece en el PPTX (puede ser un
    PPTX desactualizado con el shape de una version anterior del editor) → se
    ignora, sigue hidden, no se restaura solo
  - {hidden: true} en layout.yaml → el renderer lo omite en el deck generado

Despues de correr esto, regenera tus decks para ver los cambios.
"""
from __future__ import annotations

import colorsys
import re
import sys
from pathlib import Path

import yaml
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE as _MSO_SHAPE_TYPE

ROOT = Path(__file__).resolve().parent.parent

# Asegurar que el root del proyecto este en sys.path para importar desde src/
# (Python agrega el directorio del script a sys.path, NO el cwd, al correr scripts directamente)
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

PPTX          = ROOT / "config" / "config_visual.pptx"
GENERATED_DIR = ROOT / "config" / "generated"
LAYOUT_YAML   = GENERATED_DIR / "layout.yaml"
TEMPLATE_YAML = GENERATED_DIR / "input_template.yaml"

# Valores de muestra para los campos canonicos de cada slide type (los que el
# modelo pydantic siempre exige/acepta, independientemente de si estan ocultos
# en el editor visual). Sirven de documentacion viva de la estructura esperada.
CANONICAL_SAMPLES: dict[str, dict] = {
    "cover": {
        "type": "cover",
        "title": "Titulo Principal del Deck",
        "subtitle": "Subtitulo descriptivo de la propuesta",
        "client": "Cliente ABC S.A.",
        "date": "Junio 2026",
        "author": "Equipo Comercial",
    },
    "section_divider": {
        "type": "section_divider",
        "section_number": "01",
        "section_title": "Nombre de la Seccion",
        "tagline": "Descripcion breve de esta seccion",
    },
    "content_two_col": {
        "type": "content_two_col",
        "title": "Comparativa",
        "left": {"title": "Columna Izquierda", "bullets": ["Punto 1", "Punto 2"]},
        "right": {"title": "Columna Derecha", "bullets": ["Punto A", "Punto B"]},
    },
    "content_one_col": {
        "type": "content_one_col",
        "title": "Titulo del Slide",
        "body_title": "Encabezado de Contenido",
        "bullets": ["Primer punto del contenido", "Segundo punto del contenido"],
    },
    "pricing_table": {
        "type": "pricing_table",
        "title": "Inversion del Proyecto",
        "currency": "USD",
        "rows": [{"description": "Item de ejemplo", "quantity": 1, "unit_price": 1000.0, "total": 1000.0}],
        "totals": 1000.0,
        "notes": "Notas adicionales sobre los precios",
    },
    "profile_card": {
        "type": "profile_card",
        "name": "Nombre del Perfil",
        "role": "Rol Profesional",
        "years_experience": 5,
        "seniority": "Senior",
        "skills": ["Skill 1", "Skill 2"],
        "highlights": ["Logro destacado"],
    },
    "stat_callout": {
        "type": "stat_callout",
        "big_number": "98%",
        "label": "Descripcion del indicador clave",
        "context": "Contexto adicional del dato estadistico",
        "source": "Fuente del dato",
    },
    "closing": {
        "type": "closing",
        "headline": "Listo para transformar tu operacion?",
        "contact_name": "Nombre del Contacto",
        "contact_email": "email@softgic.com",
        "contact_phone": "+57 300 000 0000",
        "cta": "Llamado a la accion del cierre",
    },
    "grafico": {
        "type": "grafico",
        "title": "Titulo del Grafico",
        "chart_title": "Titulo del Grafico Interno",
        "chart_type": "column",
        "categories": ["Cat 1", "Cat 2"],
        "series": [{"name": "Serie 1", "values": [10, 20]}],
        "source": "Fuente del dato",
    },
}

EMU    = 914400  # EMUs por pulgada
PT_EMU = 12700   # EMUs por punto

# MSO_AUTO_SHAPE_TYPE values (confirmados con python-pptx)
_SHAPE_OVAL      = 9   # oval / ellipse
_SHAPE_RECTANGLE = 1   # rectangulo

# Solo leer shapes de estos tipos; el resto (imagenes, grupos, conectores, charts,
# tablas) se ignoran para no contaminar el layout.yaml con elementos no gestionados.
_RENDERABLE = frozenset({
    _MSO_SHAPE_TYPE.AUTO_SHAPE,  # 1 — rectangulos, ovaIos, etc.
    _MSO_SHAPE_TYPE.TEXT_BOX,    # 17 — cajas de texto
})

SLIDE_ORDER = [
    "cover", "section_divider", "content_two_col", "content_one_col",
    "pricing_table", "profile_card", "stat_callout", "closing", "grafico",
]

# Importar _D de cada renderer para saber los elementos canonicos de cada slide type.
# Esto permite marcar como hidden los elementos del _D que el usuario borro del PPTX.
try:
    from src.layout.renderers.closing import _D as _D_closing
    from src.layout.renderers.content_one_col import _D as _D_content_one_col
    from src.layout.renderers.content_two_col import _D as _D_content_two_col
    from src.layout.renderers.cover import _D as _D_cover
    from src.layout.renderers.grafico import _D as _D_grafico
    from src.layout.renderers.pricing_table import _D as _D_pricing_table
    from src.layout.renderers.profile_card import _D as _D_profile_card
    from src.layout.renderers.section_divider import _D as _D_section_divider
    from src.layout.renderers.stat_callout import _D as _D_stat_callout
    SLIDE_DEFAULTS: dict[str, dict] = {
        "cover":            _D_cover,
        "section_divider":  _D_section_divider,
        "content_two_col":  _D_content_two_col,
        "content_one_col":  _D_content_one_col,
        "pricing_table":    _D_pricing_table,
        "profile_card":     _D_profile_card,
        "stat_callout":     _D_stat_callout,
        "closing":          _D_closing,
        "grafico":          _D_grafico,
    }
except ImportError as _e:
    print(f"AVISO: no se pudo importar _D de los renderers ({_e}) — elementos borrados no se marcaran como hidden.")
    SLIDE_DEFAULTS = {}

# Campos REALES del modelo pydantic de cada slide type (declarados, no solo los
# posicionales de _D). Una forma nombrada igual a uno de estos campos (ej. 'cta'
# en closing) se resuelve dinamicamente por atributo del modelo (ver
# helpers._resolve_dynamic_value) — no es un campo "nuevo", es un campo real
# que el usuario decidio posicionar manualmente. Se usa solo para reportar esto
# con precision en la plantilla generada, distinto de un campo totalmente nuevo.
try:
    from src.models.slides import (
        ClosingSlide as _S_closing,
        ContentOneColSlide as _S_content_one_col,
        ContentTwoColSlide as _S_content_two_col,
        CoverSlide as _S_cover,
        GraficoSlide as _S_grafico,
        PricingTableSlide as _S_pricing_table,
        ProfileCardSlide as _S_profile_card,
        SectionDividerSlide as _S_section_divider,
        StatCalloutSlide as _S_stat_callout,
    )
    _SLIDE_MODEL_CLASSES = {
        "cover": _S_cover,
        "section_divider": _S_section_divider,
        "content_two_col": _S_content_two_col,
        "content_one_col": _S_content_one_col,
        "pricing_table": _S_pricing_table,
        "profile_card": _S_profile_card,
        "stat_callout": _S_stat_callout,
        "closing": _S_closing,
        "grafico": _S_grafico,
    }
    SLIDE_MODEL_FIELDS: dict[str, set[str]] = {
        slide_type: set(cls.model_fields) - {"type"}
        for slide_type, cls in _SLIDE_MODEL_CLASSES.items()
    }
    # Campos cuyo modelo pydantic los exige (sin default) — no se pueden omitir
    # del YAML de entrada aunque su elemento visual este oculto, o la validacion
    # falla con "Required field missing". Los opcionales si se pueden omitir.
    SLIDE_REQUIRED_FIELDS: dict[str, set[str]] = {
        slide_type: {
            f for f, info in cls.model_fields.items() if info.is_required() and f != "type"
        }
        for slide_type, cls in _SLIDE_MODEL_CLASSES.items()
    }
except ImportError as _e:
    print(f"AVISO: no se pudo importar los modelos pydantic ({_e}) — no se distinguiran campos reales de nuevos.")
    SLIDE_MODEL_FIELDS = {}
    SLIDE_REQUIRED_FIELDS = {}

# Mapeo campo del modelo -> elemento(s) de _D que lo usan al renderizar. Permite
# saber si un campo es relevante segun lo que el usuario tiene VISIBLE en su
# editor visual — si todos los elementos mapeados estan ocultos, el campo no
# tiene ningun efecto en el deck generado (su valor nunca se dibuja en ningun
# lado), asi que no debe pedirsele al usuario ni aparecer en el formulario web.
FIELD_TO_ELEMENTS: dict[str, dict[str, list[str]]] = {
    "cover": {
        "title": ["title"], "subtitle": ["subtitle"], "client": ["client"],
        "date": ["meta"], "author": ["meta"],
    },
    "section_divider": {
        "section_number": ["section_number"],
        "section_title": ["section_title"],
        "tagline": ["tagline"],
    },
    "content_two_col": {
        "title": ["title"],
        "left": ["left_col_title", "left_col_bullets"],
        "right": ["right_col_title", "right_col_bullets"],
    },
    "content_one_col": {
        "title": ["title"], "body_title": ["body_title"], "bullets": ["bullets"],
    },
    "pricing_table": {
        "title": ["title"], "currency": ["table"], "rows": ["table"],
        "totals": ["table"], "notes": ["notes"],
    },
    "profile_card": {
        "name": ["name"], "role": ["role"],
        "years_experience": ["seniority"], "seniority": ["seniority"],
        "skills": ["skills_area"], "photo": ["photo"],
        "highlights": ["highlights_area", "highlights_label"],
    },
    "stat_callout": {
        "big_number": ["big_number"], "label": ["label"],
        "context": ["context"], "source": ["source"],
    },
    "closing": {
        "headline": ["headline"], "contact_name": ["contact_name"],
        "contact_email": ["contact_email"], "contact_phone": ["contact_phone"],
        "cta": ["cta_box", "cta_text"],
    },
    "grafico": {
        "title": ["title"], "chart_title": ["chart_area"], "chart_type": ["chart_area"],
        "categories": ["chart_area"], "series": ["chart_area"], "source": ["source"],
    },
}

# Valores minimos validos (pasan pydantic + business_rules + schema_rules sin
# disparar nada) para campos REQUERIDOS cuando su elemento visual esta oculto.
# Cadena vacia / lista de 1 item vacio en vez del texto de muestra llamativo —
# el campo no se va a renderizar en ningun lado, asi que su contenido no importa.
MINIMAL_REQUIRED_VALUES: dict[str, dict[str, object]] = {
    "cover": {"title": ""},
    "section_divider": {"section_number": "", "section_title": ""},
    "content_two_col": {
        "title": "",
        "left": {"title": "", "bullets": [""]},
        "right": {"title": "", "bullets": [""]},
    },
    "content_one_col": {"title": "", "bullets": [""]},
    "pricing_table": {
        "title": "", "currency": "",
        "rows": [{"description": "", "quantity": 0, "unit_price": 0.0, "total": 0.0}],
        "totals": 0.0,
    },
    "profile_card": {
        "name": "", "role": "", "years_experience": 0, "seniority": "", "skills": [""],
    },
    "stat_callout": {"big_number": "", "label": ""},
    "closing": {"headline": "", "contact_name": "", "contact_email": ""},
    "grafico": {
        "title": "", "categories": [""], "series": [{"name": "", "values": [0]}],
    },
}


def is_field_visible(slide_type: str, field_name: str, section: dict) -> bool:
    """True si al menos un elemento _D mapeado a este campo esta visible
    (no oculto) en la seccion de layout.yaml. Sin mapeo conocido -> visible
    por defecto (mejor mostrar de mas que esconder algo que si importa)."""
    elements = FIELD_TO_ELEMENTS.get(slide_type, {}).get(field_name)
    if not elements:
        return True
    return any(not (isinstance(section.get(el), dict) and section[el].get("hidden")) for el in elements)


# ─────────────────────────────── utilidades ──────────────────────────────────

def _emu_to_in(v: int) -> float:
    return round(v / EMU, 3)


def _emu_to_pt(v) -> int | None:
    if v is None:
        return None
    result = round(int(v) / PT_EMU)
    return result if result > 0 else None


def _rgb_to_hex(rgb) -> str:
    # python-pptx 1.0+ RGBColor no tiene .red/.green/.blue; str() devuelve "#RRGGBB"
    s = str(rgb).lstrip('#').upper()
    if len(s) == 6:
        return s
    # Fallback para versiones antiguas que sí exponen .red/.green/.blue
    try:
        return f"{rgb.red:02X}{rgb.green:02X}{rgb.blue:02X}"
    except AttributeError:
        v = int(rgb)
        return f"{(v >> 16) & 0xFF:02X}{(v >> 8) & 0xFF:02X}{v & 0xFF:02X}"


def _get_shape_prst(shape) -> str | None:
    """Lee el 'prst' (preset geometry) directamente del XML del shape.

    Usado como fallback cuando auto_shape_type no esta en el enum de python-pptx
    (p.ej. 'snip2DiagRect', 'star5', formas de Office 2019+).
    """
    try:
        from pptx.oxml.ns import qn
        prstGeom = shape._element.spPr.find(qn('a:prstGeom'))
        return prstGeom.get('prst') if prstGeom is not None else None
    except Exception:
        return None


_RT = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships/'
_SCHEME_ALIASES = {'tx1': 'dk1', 'tx2': 'dk2', 'bg1': 'lt1', 'bg2': 'lt2'}


def _resolve_scheme_color(shape, val: str, lum_mod: int = 100000, lum_off: int = 0) -> str | None:
    """Convierte un schemeClr de PowerPoint a hex RGB navegando slide→layout→master→tema.

    Aplica luminancia modificada (lumMod, lumOff) para cubrir las variantes
    mas claro / mas oscuro del selector de colores de PowerPoint.
    """
    try:
        import lxml.etree as _etree
        from pptx.oxml.ns import qn
        elem = _SCHEME_ALIASES.get(val, val)
        if elem in ('phClr',):
            return None

        # shape puede ser un shape object (usa .part) o un slide/layout part directamente
        part   = getattr(shape, 'part', shape)
        layout = part.part_related_by(_RT + 'slideLayout')
        master = layout.part_related_by(_RT + 'slideMaster')
        theme  = master.part_related_by(_RT + 'theme')

        # theme es un Part generico — acceder via blob (XML raw) en vez de _element
        theme_xml = _etree.fromstring(theme.blob)
        scheme = theme_xml.find('.//' + qn('a:clrScheme'))
        if scheme is None:
            return None

        clr = scheme.find(qn('a:' + elem))
        if clr is None:
            return None

        srgb = clr.find(qn('a:srgbClr'))
        base = srgb.get('val', '') if srgb is not None else ''
        if not base:
            sys_c = clr.find(qn('a:sysClr'))
            base  = sys_c.get('lastClr', '') if sys_c is not None else ''
        if len(base) != 6:
            return None

        if lum_mod == 100000 and lum_off == 0:
            return base.upper()

        # Aplicar modificacion de luminancia (HLS)
        r, g, b = int(base[0:2], 16) / 255, int(base[2:4], 16) / 255, int(base[4:6], 16) / 255
        h, l, s = colorsys.rgb_to_hls(r, g, b)
        l = max(0.0, min(1.0, l * (lum_mod / 100000) + (lum_off / 100000)))
        r2, g2, b2 = colorsys.hls_to_rgb(h, l, s)
        return f'{round(r2 * 255):02X}{round(g2 * 255):02X}{round(b2 * 255):02X}'
    except Exception:
        return None


def _read_fill_hex(shape) -> str | None:
    """Lee el color de relleno solido del shape como hex RGB.

    Maneja tanto colores sRGB explicitos como schemeClr (colores de tema de PowerPoint).
    Resuelve schemeClr navegando hasta el tema del archivo PPTX.
    """
    try:
        from pptx.oxml.ns import qn
        sf = shape._element.spPr.find(qn('a:solidFill'))
        if sf is None:
            return None

        # Color sRGB explicito — leer directamente
        srgb = sf.find(qn('a:srgbClr'))
        if srgb is not None:
            v = srgb.get('val', '')
            return v.upper() if len(v) == 6 else None

        # schemeClr (color de tema) — resolver via tema del archivo
        sc = sf.find(qn('a:schemeClr'))
        if sc is not None:
            lm_e = sc.find(qn('a:lumMod'))
            lo_e = sc.find(qn('a:lumOff'))
            lm = int(lm_e.get('val', '100000')) if lm_e is not None else 100000
            lo = int(lo_e.get('val', '0'))       if lo_e is not None else 0
            return _resolve_scheme_color(shape, sc.get('val', ''), lm, lo)

    except Exception:
        pass
    return None


def _read_slide_bg(slide_pptx) -> dict | None:
    """Lee el fondo del slide del XML: soporta relleno solido y gradiente lineal.

    El usuario cambia el fondo via PowerPoint Format Background → Solid/Gradient Fill.
    Resuelve schemeClr (colores de tema) igual que _read_fill_hex para shapes.
    Devuelve {"type": "solid", "color": "RRGGBB"} o {"type": "gradient", "angle": N, "stops": [...]}
    """
    try:
        from pptx.oxml.ns import qn
        slide_part = slide_pptx.part

        bg_pr = slide_pptx._element.find('.//' + qn('p:bgPr'))
        if bg_pr is None:
            return None

        def _resolve_color(fill_elem):
            """Lee color de un fill element (solidFill o gs dentro de gradFill)."""
            srgb = fill_elem.find(qn('a:srgbClr'))
            if srgb is not None:
                v = srgb.get('val', '')
                return v.upper() if len(v) == 6 else None
            sc = fill_elem.find(qn('a:schemeClr'))
            if sc is not None:
                lm_e = sc.find(qn('a:lumMod'))
                lo_e = sc.find(qn('a:lumOff'))
                lm = int(lm_e.get('val', '100000')) if lm_e is not None else 100000
                lo = int(lo_e.get('val', '0'))       if lo_e is not None else 0
                return _resolve_scheme_color(slide_part, sc.get('val', ''), lm, lo)
            return None

        # ── Relleno solido ───────────────────────────────────────────────────
        sf = bg_pr.find(qn('a:solidFill'))
        if sf is not None:
            color = _resolve_color(sf)
            if color:
                return {"type": "solid", "color": color}

        # ── Gradiente lineal ─────────────────────────────────────────────────
        gf = bg_pr.find(qn('a:gradFill'))
        if gf is None:
            return None
        gs_lst = gf.find(qn('a:gsLst'))
        if gs_lst is None:
            return None

        stops = []
        for gs in gs_lst.findall(qn('a:gs')):
            pos = round(int(gs.get('pos', '0')) / 1000)  # 0-100000 → 0-100
            color = _resolve_color(gs)
            if color:
                stops.append({"pos": pos, "color": color})

        if not stops:
            return None

        lin = gf.find(qn('a:lin'))
        angle = round(int(lin.get('ang', '0')) / 60000) if lin is not None else 90

        return {"type": "gradient", "angle": angle, "stops": stops}

    except Exception:
        return None


def _read_autofit_scale(shape) -> float:
    """Lee el factor de escala que PowerPoint aplica via 'Reducir texto al
    desbordar' (normAutofit/@fontScale). El tamano de fuente en el XML del run
    es el NOMINAL configurado por el usuario — si el texto no entraba en la
    forma, PowerPoint lo encoge visualmente sin tocar ese valor nominal. Sin
    aplicar este factor, el texto se regenera mas grande de lo que realmente
    se veia y no encaja en las mismas coordenadas/proporcion.
    """
    try:
        from pptx.oxml.ns import qn
        bodyPr = shape.text_frame._txBody.find(qn('a:bodyPr'))
        if bodyPr is None:
            return 1.0
        norm_autofit = bodyPr.find(qn('a:normAutofit'))
        if norm_autofit is None:
            return 1.0
        scale = norm_autofit.get('fontScale')
        return int(scale) / 100000 if scale else 1.0
    except Exception:
        return 1.0


def _read_text_margins(shape) -> dict:
    """Lee los margenes internos (insets) del text_frame en pulgadas. Sin
    replicar esto, el ancho/alto util para el texto difiere del original y el
    wrap/tamano visual ya no encaja en la misma caja.
    """
    try:
        tf = shape.text_frame
        return {
            "margin_left": round((tf.margin_left or 0) / EMU, 3),
            "margin_top": round((tf.margin_top or 0) / EMU, 3),
            "margin_right": round((tf.margin_right or 0) / EMU, 3),
            "margin_bottom": round((tf.margin_bottom or 0) / EMU, 3),
        }
    except Exception:
        return {}


def _read_run_font_size(run) -> int | None:
    """Lee el tamano de fuente del run: API python-pptx y fallback via XML sz."""
    try:
        sz_emu = run.font.size
        if sz_emu is not None:
            result = round(int(sz_emu) / 12700)
            return result if result > 0 else None
    except Exception:
        pass
    try:
        from pptx.oxml.ns import qn
        rPr = run._r.find(qn('a:rPr'))
        if rPr is not None:
            sz_str = rPr.get('sz')
            if sz_str:
                return max(1, int(sz_str) // 100)
    except Exception:
        pass
    return None


def _read_para_default_font_size(para) -> int | None:
    """Lee el tamano de fuente por defecto del parrafo via pPr/defRPr/sz."""
    try:
        from pptx.oxml.ns import qn
        pPr = para._p.find(qn('a:pPr'))
        if pPr is not None:
            defRPr = pPr.find(qn('a:defRPr'))
            if defRPr is not None:
                sz_str = defRPr.get('sz')
                if sz_str:
                    return max(1, int(sz_str) // 100)
    except Exception:
        pass
    return None


def _read_run_color_hex(run, shape) -> str | None:
    """Lee el color del texto del run: API para sRGB, XML para colores de tema."""
    # Primero intentar via API (rapido, funciona para colores sRGB explicitos)
    try:
        clr = run.font.color
        if clr and clr.type is not None:
            return _rgb_to_hex(clr.rgb)
    except Exception:
        pass

    # Fallback via XML para schemeClr (colores del tema de PowerPoint)
    try:
        from pptx.oxml.ns import qn
        rPr = run._r.find(qn('a:rPr'))
        if rPr is None:
            return None
        sf = rPr.find(qn('a:solidFill'))
        if sf is None:
            return None

        srgb = sf.find(qn('a:srgbClr'))
        if srgb is not None:
            v = srgb.get('val', '')
            return v.upper() if len(v) == 6 else None

        sc = sf.find(qn('a:schemeClr'))
        if sc is not None:
            lm_e = sc.find(qn('a:lumMod'))
            lo_e = sc.find(qn('a:lumOff'))
            lm = int(lm_e.get('val', '100000')) if lm_e is not None else 100000
            lo = int(lo_e.get('val', '0'))       if lo_e is not None else 0
            return _resolve_scheme_color(shape, sc.get('val', ''), lm, lo)
    except Exception:
        pass
    return None


def _extract_text_content(shape) -> dict:
    """Extrae texto/runs/font_size/bold/color de cualquier shape con text_frame
    (TEXT_BOX o un AUTO_SHAPE usado como contenedor de texto). Devuelve {} si
    el shape no tiene texto.
    """
    result: dict = {}
    try:
        if not shape.has_text_frame:
            return result
    except Exception:
        return result

    # Factor de "Reducir texto al desbordar" — el tamano nominal en el XML no
    # es el tamano realmente renderizado si PowerPoint lo encogio para que
    # quepa en la forma. Se aplica a cada font_size leido mas abajo.
    autofit_scale = _read_autofit_scale(shape)

    def _scaled(size: int | None) -> int | None:
        if size is None or autofit_scale == 1.0:
            return size
        return max(1, round(size * autofit_scale))

    # Iterar TODOS los runs de todos los parrafos para capturar font_size,
    # bold, color principal, y datos por-run para texto multicolor.
    all_run_data: list[dict] = []
    seen_colors: set[str] = set()

    for para in shape.text_frame.paragraphs:
        for run in para.runs:
            run_text = run.text
            run_size = _scaled(_read_run_font_size(run))
            run_bold = run.font.bold
            run_color = _read_run_color_hex(run, shape)

            # Propiedades a nivel shape (primera encontrada)
            if "font_size" not in result and run_size:
                result["font_size"] = run_size
            if "bold" not in result and run_bold is not None:
                result["bold"] = run_bold
            if "color" not in result and run_color:
                result["color"] = run_color

            if run_text:
                run_info: dict = {"text": run_text}
                if run_color:
                    run_info["color"] = run_color
                    seen_colors.add(run_color)
                if run_bold is not None:
                    run_info["bold"] = run_bold
                if run_size:
                    run_info["font_size"] = run_size
                all_run_data.append(run_info)

        # Fallback: tamano por defecto del parrafo si todavia no se encontro
        if "font_size" not in result:
            sz = _scaled(_read_para_default_font_size(para))
            if sz:
                result["font_size"] = sz

    # Guardar runs solo si hay multiples colores distintos (texto multicolor)
    if len(seen_colors) > 1 and 0 < len(all_run_data) <= 30:
        result["runs"] = all_run_data

    try:
        text = shape.text_frame.text.strip()
        if text:
            result["text"] = text
    except Exception:
        pass

    # Margenes internos: solo importan si efectivamente hay texto que ajustar.
    if result.get("text"):
        result.update(_read_text_margins(shape))

    return result


def _read_shape_props(shape) -> dict:
    """Lee posicion, kind (box|oval|shape|text), color de relleno y estilo de
    fuente. Solo debe recibir shapes de tipo AUTO_SHAPE o TEXT_BOX (filtrado
    en main loop).
    """
    props: dict = {
        "left":   _emu_to_in(shape.left),
        "top":    _emu_to_in(shape.top),
        "width":  _emu_to_in(shape.width),
        "height": _emu_to_in(shape.height),
    }

    # Rotacion aplicada en PowerPoint (Formato de forma -> Tamano -> Giro) — sin
    # capturarla, una forma girada vuelve a renderizarse en 0 grados, perdiendo
    # su orientacion visual. round() evita arrastrar ruido de punto flotante.
    try:
        rot = round(float(shape.rotation), 1)
        if rot:
            props["rotation"] = rot
    except Exception:
        pass

    try:
        if shape.shape_type == _MSO_SHAPE_TYPE.TEXT_BOX:
            props["kind"] = "text"
            props.update(_extract_text_content(shape))

        else:
            # AUTO_SHAPE — detectar tipo especifico

            # Intento 1: leer auto_shape_type del enum de python-pptx
            ast: int | None = None
            try:
                raw = shape.auto_shape_type
                if raw is not None:
                    ast = int(raw)
            except (AttributeError, TypeError, KeyError, ValueError):
                # python-pptx 1.0.x lanza ValueError (no None) cuando el 'prst'
                # del shape no esta en su enum MSO_AUTO_SHAPE_TYPE (formas nuevas
                # de PowerPoint 2016+ sin mapeo). Sin capturar esto, el ValueError
                # se propagaba al except exterior y descartaba relleno/texto del
                # shape entero, degradandolo a una caja vacia sin contenido.
                pass

            if ast == _SHAPE_OVAL:
                props["kind"] = "oval"
            elif ast == _SHAPE_RECTANGLE or ast is None and _get_shape_prst(shape) in (None, "rect"):
                props["kind"] = "box"
            elif ast is not None:
                # Tipo en el enum pero no es ovalo ni rectangulo (estrella, flecha, etc.)
                props["kind"] = "shape"
                props["shape_type_id"] = ast
            else:
                # Intento 2: el prst no esta en el enum de python-pptx
                # (formas de Office 2019+, snip corners, etc.) — leer desde XML
                prst = _get_shape_prst(shape)
                if prst == "ellipse":
                    props["kind"] = "oval"
                elif prst:
                    props["kind"] = "shape"
                    props["shape_prst"] = prst
                else:
                    props["kind"] = "box"

            # Leer color de relleno solido: soporta sRGB y colores de tema
            try:
                if shape.fill.type == 1:  # SOLID
                    hex_val = _read_fill_hex(shape)
                    if hex_val:
                        props["fill"] = hex_val
            except Exception:
                pass

            # Muchas formas (rectangulos, etc.) se usan como CONTENEDORES DE TEXTO
            # sin relleno visible, en vez de cajas de texto puras — sin esto, su
            # texto se descartaba silenciosamente y la forma se renderizaba como
            # un bloque de color solido vacio (el fill por defecto del slide).
            text_content = _extract_text_content(shape)
            if text_content.get("text"):
                if "fill" not in props:
                    # Sin relleno visible -> es un contenedor de texto puro,
                    # no una forma decorativa — descartar kind/shape_* previos
                    # (pero conservar rotation, ya leida arriba).
                    props = {
                        "left": props["left"], "top": props["top"],
                        "width": props["width"], "height": props["height"],
                        "kind": "text",
                        **({"rotation": props["rotation"]} if "rotation" in props else {}),
                    }
                # Con relleno: se conserva el kind (box/oval/shape) Y el texto,
                # para que el deck dibuje la forma de color CON su texto encima
                # (ej. un boton/CTA solido con una etiqueta).
                props.update(text_content)

    except Exception:
        props.setdefault("kind", "box")

    return props


def _safe_filename(name: str, ext: str) -> str:
    """Convierte nombre de shape a nombre de archivo seguro."""
    safe = re.sub(r'[^\w\-.]', '_', name).strip('_').lower()
    if not safe:
        safe = "image"
    ext = ext.lstrip('.').lower()
    # Normalizar jpeg → jpg
    ext = 'jpg' if ext == 'jpeg' else ext
    if not safe.endswith(f".{ext}"):
        safe = f"{safe}.{ext}"
    return safe


def _read_picture_props(shape, slide_type: str) -> dict | None:
    """Extrae posicion e imagen embebida de un shape PICTURE.

    Guarda los bytes de la imagen en config/images/<slide_type>/ y retorna
    props con kind:image y src relativo al root del proyecto.
    """
    try:
        img = shape.image
    except AttributeError:
        return None

    ext = img.ext or "png"
    filename = _safe_filename(shape.name, ext)

    images_dir = ROOT / "config" / "images" / slide_type
    images_dir.mkdir(parents=True, exist_ok=True)

    dest = images_dir / filename
    dest.write_bytes(img.blob)

    return {
        "left":   _emu_to_in(shape.left),
        "top":    _emu_to_in(shape.top),
        "width":  _emu_to_in(shape.width),
        "height": _emu_to_in(shape.height),
        "kind":   "image",
        "src":    f"config/images/{slide_type}/{filename}",
    }


def _props_changed(old: dict, new: dict) -> list[str]:
    """Retorna lista de claves que cambiaron (ignorando solo 'hidden')."""
    keys = {k for k in (set(old) | set(new)) if k != "hidden"}
    return [k for k in keys if old.get(k) != new.get(k)]


# "type" es el discriminador del modelo pydantic (cover, closing, etc.) — una
# forma nombrada asi NUNCA debe tratarse como campo dinamico, o corromperia la
# entrada del deck al sobreescribir el tipo de slide en la plantilla generada.
_RESERVED_DYNAMIC_NAMES = frozenset({"type"})


def _dynamic_text_fields(section: dict, defaults_keys: set[str]) -> dict[str, str]:
    """Detecta formas de texto o imagen NUEVAS (no canonicas) nombradas por el
    usuario en el editor visual. Cada una se vuelve un campo dinamico: su nombre
    es la clave que debe usarse en el YAML de entrada para que el deck muestre
    un valor real ahi (texto literal, o ruta de imagen para kind: image).

    Incluye tambien formas de color (kind: box/oval/shape) que ADEMAS tienen
    texto encima (ej. un boton/CTA solido con etiqueta) — sin esto, su texto
    quedaba capturado en layout.yaml pero invisible/no-editable en la
    plantilla y el formulario web, aunque el deck si lo dibujara.
    """
    fields: dict[str, str] = {}
    for name, props in section.items():
        if name.startswith("_") or name in defaults_keys or name == "background_color":
            continue
        if name in _RESERVED_DYNAMIC_NAMES:
            continue
        if not isinstance(props, dict) or props.get("hidden"):
            continue
        kind = props.get("kind")
        if kind == "image":
            fields[name] = props.get("src") or f"ruta/a/imagen_{name}.png"
        elif props.get("text"):
            # kind == "text", o una forma de color con texto encima
            fields[name] = props["text"]
    return fields


def _write_input_template(current: dict) -> None:
    """Genera config/generated/input_template.yaml: un YAML de ejemplo con las 9 slides,
    listo para copiar y llenar con datos reales. Incluye los campos canonicos
    de cada slide type mas cualquier campo dinamico detectado en layout.yaml.

    Distingue dos casos de formas no-canonicas nombradas en el editor visual:
    - Campo NUEVO (no existe en el modelo pydantic): se agrega como clave nueva
      a la plantilla con el texto de muestra capturado.
    - Campo REAL reposicionado (ej. nombrar una forma 'cta' en closing, que ya
      es un campo del modelo pero no esta en _D): NO se agrega como clave nueva
      (ya esta en CANONICAL_SAMPLES) — solo se informa en el encabezado, para no
      confundirlo con un campo inventado.
    """
    new_fields_by_slide: dict[str, dict[str, str]] = {}
    bound_fields_by_slide: dict[str, list[str]] = {}
    omitted_by_slide: dict[str, list[str]] = {}
    slides_out: list[dict] = []

    for slide_type in SLIDE_ORDER:
        section = current.get(slide_type, {})
        if not isinstance(section, dict):
            section = {}
        defaults_keys = set(SLIDE_DEFAULTS.get(slide_type, {}).keys())
        model_fields = SLIDE_MODEL_FIELDS.get(slide_type, set())
        required_fields = SLIDE_REQUIRED_FIELDS.get(slide_type, set())
        detected = _dynamic_text_fields(section, defaults_keys)

        new_fields = {k: v for k, v in detected.items() if k not in model_fields}
        bound_fields = [k for k in detected if k in model_fields]

        # Solo incluir campos canonicos cuyo elemento visual SI esta presente
        # en el diseño del usuario. Si esta oculto (reemplazado por una forma
        # propia, o borrado): los opcionales se omiten del todo; los requeridos
        # por el modelo se mantienen pero con un valor minimo neutro (no se
        # puede omitir un campo requerido sin romper la validacion).
        sample: dict = {"type": slide_type}
        omitted: list[str] = []
        for field_name, sample_value in CANONICAL_SAMPLES.get(slide_type, {}).items():
            if field_name == "type":
                continue
            if is_field_visible(slide_type, field_name, section):
                sample[field_name] = sample_value
            elif field_name in required_fields:
                minimal = MINIMAL_REQUIRED_VALUES.get(slide_type, {}).get(field_name, sample_value)
                sample[field_name] = minimal
            else:
                omitted.append(field_name)

        sample.update(new_fields)
        slides_out.append(sample)

        if new_fields:
            new_fields_by_slide[slide_type] = new_fields
        if bound_fields:
            bound_fields_by_slide[slide_type] = bound_fields
        if omitted:
            omitted_by_slide[slide_type] = omitted

    header_lines = [
        "# config/generated/input_template.yaml — generado automaticamente por 'Aplicar cambios del editor visual'",
        "# Copia este archivo, renombralo y reemplaza los valores de ejemplo con datos reales.",
        "# Solo incluye campos cuyo elemento esta VISIBLE en tu editor visual actual.",
    ]
    if omitted_by_slide:
        header_lines.append("#")
        header_lines.append("# Campos omitidos (su elemento esta oculto/reemplazado en tu diseño,")
        header_lines.append("# no tienen ningun efecto en el deck generado):")
        for slide_type, fields in omitted_by_slide.items():
            header_lines.append(f"#   {slide_type}: {', '.join(fields)}")
    if new_fields_by_slide:
        header_lines.append("#")
        header_lines.append("# Campos dinamicos NUEVOS detectados (formas sin campo correspondiente en")
        header_lines.append("# el modelo — ya estan agregados abajo con su valor de muestra):")
        for slide_type, fields in new_fields_by_slide.items():
            header_lines.append(f"#   {slide_type}: {', '.join(fields.keys())}")
    if bound_fields_by_slide:
        header_lines.append("#")
        header_lines.append("# Campos existentes reposicionados manualmente en el editor visual")
        header_lines.append("# (ya forman parte del modelo, solo cambiaron de estilo/posicion):")
        for slide_type, fields in bound_fields_by_slide.items():
            header_lines.append(f"#   {slide_type}: {', '.join(fields)}")

    content = (
        "\n".join(header_lines) + "\n\n"
        + yaml.dump({"slides": slides_out}, default_flow_style=False, allow_unicode=True, sort_keys=False)
    )
    tmp = TEMPLATE_YAML.with_suffix(".yaml.tmp")
    try:
        tmp.write_text(content, encoding="utf-8")
        tmp.replace(TEMPLATE_YAML)
    except Exception:
        TEMPLATE_YAML.write_text(content, encoding="utf-8")
        try:
            tmp.unlink(missing_ok=True)
        except Exception:
            pass

    print(f"\nPlantilla de datos generada: {TEMPLATE_YAML}")
    if omitted_by_slide:
        total_omitted = sum(len(v) for v in omitted_by_slide.values())
        print(f"  ({total_omitted} campo(s) omitido(s) por estar oculto(s) en {len(omitted_by_slide)} slide(s))")
    if new_fields_by_slide:
        total = sum(len(v) for v in new_fields_by_slide.values())
        print(f"  ({total} campo(s) dinamico(s) nuevo(s) en {len(new_fields_by_slide)} slide(s))")
    if bound_fields_by_slide:
        total_bound = sum(len(v) for v in bound_fields_by_slide.values())
        print(f"  ({total_bound} campo(s) existente(s) reposicionado(s) en {len(bound_fields_by_slide)} slide(s))")


# ─────────────────────────────────── main ────────────────────────────────────

def main() -> None:
    if not PPTX.exists():
        print(f"ERROR: No se encontro {PPTX}")
        print("Primero ejecuta: Menu -> [3] Herramientas de diseno -> [1] Abrir editor visual")
        sys.exit(1)

    GENERATED_DIR.mkdir(parents=True, exist_ok=True)

    prs = Presentation(str(PPTX))

    if len(prs.slides) != len(SLIDE_ORDER):
        print(f"ERROR: Se esperaban {len(SLIDE_ORDER)} slides, el archivo tiene {len(prs.slides)}")
        print("El archivo puede estar desactualizado. Regenera el editor visual.")
        sys.exit(1)

    current: dict = {}
    if LAYOUT_YAML.exists():
        try:
            loaded = yaml.safe_load(LAYOUT_YAML.read_text(encoding="utf-8"))
            current = loaded if isinstance(loaded, dict) else {}
        except Exception as exc:
            print(f"AVISO: no se pudo leer layout.yaml existente ({exc}) — se generara desde cero")
            current = {}

    updated = 0
    for slide_pptx, slide_type in zip(prs.slides, SLIDE_ORDER):
        old_section   = current.get(slide_type, {})
        if not isinstance(old_section, dict):
            old_section = {}
        defaults_keys = set(SLIDE_DEFAULTS.get(slide_type, {}).keys())

        # ── Leer shapes presentes en el PPTX ────────────────────────────────
        new_section: dict = {}

        # Leer color/gradiente de fondo del slide (Format Background en PowerPoint)
        bg_data = _read_slide_bg(slide_pptx)
        if bg_data:
            new_section["_background"] = bg_data

        for shape in slide_pptx.shapes:
            name = shape.name
            if name.startswith("_"):
                continue  # shapes internos del editor (ej: _LABEL_)

            if name in new_section:
                print(f"  AVISO [{slide_type}]: nombre duplicado '{name}' — se toma la ultima instancia")

            # Elemento legacy del editor visual anterior (indicador de color de fondo)
            if name == "background_color":
                continue

            # Elemento canonico previamente marcado como hidden por el usuario:
            # NO restaurar aunque este en el PPTX (puede ser un PPTX desactualizado
            # que todavia tiene el shape del editor visual anterior). Se aplica
            # ANTES de la rama de PICTURE para cubrir tambien canonicos de tipo
            # imagen (ej. 'logo', 'photo') — si no, una imagen desactualizada
            # con ese nombre podia restaurar un elemento que el usuario borro.
            if name in defaults_keys and isinstance(old_section.get(name), dict) and old_section[name].get("hidden"):
                continue

            # Imagenes embebidas: extraer bytes y guardar en config/images/
            if shape.shape_type == _MSO_SHAPE_TYPE.PICTURE:
                try:
                    props = _read_picture_props(shape, slide_type)
                    if props:
                        new_section[name] = props
                        print(f"  [{slide_type}] imagen '{name}' -> {props['src']}")
                except Exception as exc:
                    print(f"  AVISO [{slide_type}]: no se pudo leer imagen '{name}': {exc}")
                continue

            # Ignorar tipos de shape no gestionados: grupos, conectores, charts, tablas.
            if shape.shape_type not in _RENDERABLE:
                continue

            try:
                props = _read_shape_props(shape)
                # Para elementos canonicos no guardar el texto de muestra del editor visual
                if name in defaults_keys:
                    props.pop("text", None)
                new_section[name] = props
            except Exception as exc:
                print(f"  AVISO [{slide_type}]: no se pudo leer '{name}': {exc}")

        # ── Construir seccion final ──────────────────────────────────────────
        # 1. Empezar con lo que esta en el PPTX (posicion y kind actualizados)
        final_section: dict = dict(new_section)

        # 2. Elementos de _D ausentes del PPTX → ocultos en el deck. El editor visual
        # es la fuente de verdad: si el elemento no esta en el PPTX (lo borraste,
        # o nunca lo agregaste), no debe aparecer en el deck generado, sin importar
        # el historial en old_section.
        # (El caso de un PPTX desactualizado que aun tiene el shape se maneja arriba,
        # en el bucle de lectura de shapes, que ignora elementos previamente hidden.)
        for k in defaults_keys:
            if k not in new_section:
                final_section[k] = {"hidden": True}

        # 3. Extras (no en _D) que estaban en old_section pero ya no en PPTX
        #    → hidden (formas normales) o conservar (claves internas _*)
        for k in old_section:
            if k not in new_section and k not in defaults_keys:
                if k.startswith("_"):
                    final_section[k] = old_section[k]  # conservar valor interno si no lo detectamos
                else:
                    final_section[k] = {"hidden": True}

        # ── Detectar y reportar cambios ──────────────────────────────────────
        added: list[str] = []
        for k in new_section:
            if k.startswith("_"):
                continue  # no reportar claves internas como "nuevo/restaurado"
            was_hidden = isinstance(old_section.get(k), dict) and old_section[k].get("hidden")
            if was_hidden:
                added.append(f"  [{slide_type}] +{k} (restaurado)")
            elif k not in old_section:
                added.append(f"  [{slide_type}] +{k} (nuevo)")

        newly_hidden: list[str] = [
            k for k, v in final_section.items()
            if not k.startswith("_")
            and v == {"hidden": True}
            and not (isinstance(old_section.get(k), dict) and old_section[k].get("hidden"))
        ]

        changed_lines: list[str] = []
        for k in new_section:
            if k.startswith("_"):
                continue  # no reportar cambios en claves internas
            old_k = old_section.get(k)
            if isinstance(old_k, dict) and not old_k.get("hidden"):
                for prop in _props_changed(old_k, new_section[k]):
                    changed_lines.append(
                        f"  [{slide_type}] {k}.{prop}: {old_k.get(prop)} -> {new_section[k].get(prop)}"
                    )

        if added or newly_hidden or changed_lines:
            for msg in added:
                print(msg)
            for k in newly_hidden:
                print(f"  [{slide_type}] -{k} (eliminado)")
            for line in changed_lines:
                print(line)
            updated += 1

        current[slide_type] = final_section

    # ── Escritura atomica del YAML ───────────────────────────────────────────
    # Escribir primero a archivo temporal y luego renombrar para evitar
    # corrupcion si el proceso se interrumpe a mitad de la escritura.
    content = (
        "# config/generated/layout.yaml — actualizado desde editor visual\n"
        "# Medidas en pulgadas (1\" = 2.54 cm). font_size en puntos. color/fill en hex RGB.\n\n"
        + yaml.dump(current, default_flow_style=False, allow_unicode=True, sort_keys=False)
    )
    tmp = LAYOUT_YAML.with_suffix(".yaml.tmp")
    try:
        tmp.write_text(content, encoding="utf-8")
        tmp.replace(LAYOUT_YAML)  # atomico: falla o tiene exito completo
    except Exception:
        # Fallback no-atomico si replace() falla (permisos, cross-device, etc.)
        LAYOUT_YAML.write_text(content, encoding="utf-8")
        try:
            tmp.unlink(missing_ok=True)
        except Exception:
            pass

    # Limpiar cache de layout_config si esta cargado en este proceso
    try:
        from src.layout import layout_config
        layout_config.reload()
    except Exception:
        pass

    if updated == 0:
        print("Sin cambios detectados.")
    else:
        print(f"\n{updated} tipo(s) de slide actualizados en config/generated/layout.yaml")
        print("Regenera tus decks para ver los cambios.")

    _write_input_template(current)


if __name__ == "__main__":
    main()
