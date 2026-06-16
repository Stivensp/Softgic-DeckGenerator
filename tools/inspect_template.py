"""
Lee assets/template.pptx y muestra todas las posiciones de shapes
que ya existen en el Slide Master y en cada layout.

Usa esto cuando recibes un template maestro de Softgic para saber
exactamente donde estan posicionados los elementos del diseno.

Medidas en pulgadas (in) y centimetros (cm).
"""
from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.util import Emu

ROOT = Path(__file__).resolve().parent.parent

EMU_PER_INCH = 914400
EMU_PER_CM   = 360000


def emu_to_in(emu: int) -> float:
    return round(emu / EMU_PER_INCH, 3)


def emu_to_cm(emu: int) -> float:
    return round(emu / EMU_PER_CM, 2)


def print_shape(shape, indent: str = "    ") -> None:
    try:
        left   = emu_to_in(shape.left)
        top    = emu_to_in(shape.top)
        width  = emu_to_in(shape.width)
        height = emu_to_in(shape.height)
        l_cm   = emu_to_cm(shape.left)
        t_cm   = emu_to_cm(shape.top)
        w_cm   = emu_to_cm(shape.width)
        h_cm   = emu_to_cm(shape.height)
    except Exception:
        print(f"{indent}[shape sin posicion: {shape.name}]")
        return

    print(f"{indent}{shape.name}")
    print(f"{indent}  pulgadas -> left={left}  top={top}  width={width}  height={height}")
    print(f"{indent}  cm       -> left={l_cm}  top={t_cm}  width={w_cm}  height={h_cm}")

    # Si tiene texto, mostrarlo
    try:
        if shape.has_text_frame and shape.text_frame.text.strip():
            txt = shape.text_frame.text.strip()[:60]
            print(f"{indent}  texto: \"{txt}\"")
    except Exception:
        pass


def main() -> None:
    template = ROOT / "assets" / "template.pptx"
    if not template.exists():
        print(f"ERROR: No se encontro {template}")
        return

    prs = Presentation(str(template))

    w_in = emu_to_in(prs.slide_width)
    h_in = emu_to_in(prs.slide_height)
    w_cm = emu_to_cm(prs.slide_width)
    h_cm = emu_to_cm(prs.slide_height)

    print("=" * 60)
    print(f"  INSPECCION DE TEMPLATE: {template.name}")
    print("=" * 60)
    print(f"  Dimensiones: {w_in}\" x {h_in}\"  ({w_cm}cm x {h_cm}cm)")
    print(f"  Layouts disponibles: {len(prs.slide_layouts)}")
    print()

    # Slide Master
    master = prs.slide_master
    master_shapes = list(master.shapes)
    if master_shapes:
        print(f"SLIDE MASTER ({len(master_shapes)} shapes)")
        print("-" * 40)
        for shape in master_shapes:
            print_shape(shape)
        print()

    # Layouts
    for i, layout in enumerate(prs.slide_layouts):
        shapes = list(layout.shapes)
        print(f"LAYOUT [{i}]: {layout.name}  ({len(shapes)} shapes)")
        print("-" * 40)
        if shapes:
            for shape in shapes:
                print_shape(shape)
        else:
            print("    (sin shapes propios — hereda del master)")
        print()

    print("=" * 60)
    print("  COMO USAR ESTA INFORMACION")
    print("=" * 60)
    print()
    print("  Si el template tiene shapes en el Slide Master o en un Layout,")
    print("  sus posiciones son las que debes usar en los renderers.")
    print()
    print("  Ejemplo: si el logo del master esta en left=0.5 top=0.25,")
    print("  busca en src/layout/renderers/cover.py la llamada a add_image()")
    print("  y ajusta left_in=0.5, top_in=0.25")
    print()
    print("  Si el template no tiene shapes (solo layouts vacios),")
    print("  usa el archivo output/PREVIEW_layouts.pptx como referencia visual.")
    print("  Ejecuta: python tools/preview_layout.py")


if __name__ == "__main__":
    main()
