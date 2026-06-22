"""
Genera config/config_visual.pptx y lo abre en PowerPoint.

Cada slide muestra el diseno REAL del tipo de layout:
- Fondo del color correcto (navy, blanco, etc.)
- Texto de ejemplo con la fuente, tamano y color reales
- Barras y paneles con sus colores de tema reales

COMO USAR:
  1. Menu -> [3] Herramientas de diseno -> [1] Abrir editor visual
  2. PowerPoint se abre mostrando como quedara cada slide
  3. Mueve, redimensiona o cambia colores/fuentes a tu gusto
  4. Guarda (Ctrl+S) y cierra PowerPoint
  5. Menu -> [3] -> [2] Aplicar cambios  (lee posiciones, tamanoss y colores)
"""
from __future__ import annotations

import os
import sys
from pathlib import Path

import yaml
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parent.parent
LAYOUT_YAML = ROOT / "config" / "generated" / "layout.yaml"
THEME_YAML = ROOT / "config" / "theme.yaml"
OUT_PPTX = ROOT / "config" / "config_visual.pptx"

W, H = 13.33, 7.5
_RECT = 1  # MSO_AUTO_SHAPE_TYPE.RECTANGLE
_OVAL = 9  # MSO_AUTO_SHAPE_TYPE.OVAL


def _rgb(hex_color: str) -> RGBColor:
    h = hex_color.lstrip("#")
    return RGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))


def _load_theme() -> dict:
    if THEME_YAML.exists():
        return yaml.safe_load(THEME_YAML.read_text(encoding="utf-8")) or {}
    return {}


def _label_color(bg_hex: str) -> RGBColor:
    """White label on dark backgrounds, dark on light."""
    h = bg_hex.lstrip("#")
    r, g, b = int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16)
    luminance = 0.299 * r + 0.587 * g + 0.114 * b
    return RGBColor(255, 255, 255) if luminance < 140 else RGBColor(30, 30, 30)


# ──────────────────────────────────────────────────────────────
# Definicion de cada elemento por slide type
# t: "text" | "fill" | "image"
# fill / text elements se stylizan con los valores del tema
# ──────────────────────────────────────────────────────────────
def _build_defs(c: dict, f: dict) -> dict:
    """Construye las definiciones de diseno usando los colores/fuentes del tema."""
    return {
        "cover": {
            "_bg": c["background"],
            "accent_bar":  {"t": "fill", "fill": c["accent"]},
            "logo":        {"t": "image", "label": "LOGOTIPO"},
            "title":       {"t": "text", "text": "Titulo Principal del Deck", "pt": f["size_heading"] + 4, "color": c["text_dark"], "bold": True},
            "subtitle":    {"t": "text", "text": "Subtitulo descriptivo de la propuesta", "pt": f["size_subheading"] - 4, "color": c["accent"]},
            "divider":     {"t": "fill", "fill": c["accent"]},
            "client":      {"t": "text", "text": "CLIENTE ABC S.A.", "pt": f["size_body"], "color": c["text_muted"]},
            "meta":        {"t": "text", "text": "Autor  |  Fecha", "pt": f["size_caption"], "color": c["text_muted"], "align": "right"},
        },
        "section_divider": {
            "_bg": c["background"],
            "top_line":       {"t": "fill", "fill": c["accent"]},
            "right_bar":      {"t": "fill", "fill": c["accent"]},
            "section_number": {"t": "text", "text": "01", "pt": f["size_stat"], "color": c["accent"], "bold": True},
            "section_title":  {"t": "text", "text": "Nombre de la Seccion", "pt": f["size_heading"], "color": c["text_dark"], "bold": True},
            "divider":        {"t": "fill", "fill": c["accent"]},
            "tagline":        {"t": "text", "text": "Descripcion breve de esta seccion", "pt": f["size_body"], "color": c["text_muted"], "italic": True},
        },
        "content_two_col": {
            "_bg": c["background"],
            "header_bar":        {"t": "fill", "fill": c["primary"]},
            "title":             {"t": "text", "text": "Titulo del Slide de Contenido", "pt": f["size_subheading"], "color": c["text_light"], "bold": True},
            "accent_line":       {"t": "fill", "fill": c["accent"]},
            "left_col_title":    {"t": "text", "text": "Columna Izquierda", "pt": f["size_body"] + 1, "color": c["accent"], "bold": True},
            "left_col_bullets":  {"t": "text", "text": "Punto 1\nPunto 2\nPunto 3", "pt": f["size_body"], "color": c["text_dark"]},
            "divider":           {"t": "fill", "fill": c["divider"]},
            "right_col_title":   {"t": "text", "text": "Columna Derecha", "pt": f["size_body"] + 1, "color": c["accent"], "bold": True},
            "right_col_bullets": {"t": "text", "text": "Punto A\nPunto B\nPunto C", "pt": f["size_body"], "color": c["text_dark"]},
        },
        "content_one_col": {
            "_bg": c["background"],
            "header_bar":  {"t": "fill", "fill": c["primary"]},
            "title":       {"t": "text", "text": "Titulo del Slide", "pt": f["size_subheading"], "color": c["text_light"], "bold": True},
            "accent_line": {"t": "fill", "fill": c["accent"]},
            "body_title":  {"t": "text", "text": "Encabezado de Contenido", "pt": f["size_body"] + 2, "color": c["primary"], "bold": True},
            "bullets":     {"t": "text", "text": "Primer punto del contenido\nSegundo punto del contenido\nTercer punto del contenido", "pt": f["size_body"], "color": c["text_dark"]},
        },
        "pricing_table": {
            "_bg": c["background"],
            "header_bar":  {"t": "fill", "fill": c["primary"]},
            "title":       {"t": "text", "text": "Titulo de la Tabla de Precios", "pt": f["size_subheading"], "color": c["text_light"], "bold": True},
            "accent_line": {"t": "fill", "fill": c["accent"]},
            "table":       {"t": "fill", "fill": c["divider"], "label": "TABLA DE PRECIOS (Descripcion | Cant | Unidad | P.Unit | Total)"},
            "notes":       {"t": "text", "text": "* Notas adicionales sobre los precios", "pt": f["size_caption"], "color": c["text_muted"], "italic": True},
        },
        "profile_card": {
            "_bg": c["background"],
            "left_panel":      {"t": "fill", "fill": c["primary"]},
            "left_accent_bar": {"t": "fill", "fill": c["accent"]},
            "photo":           {"t": "image", "label": "FOTO\nPERFIL"},
            "name":            {"t": "text", "text": "Nombre del Perfil", "pt": f["size_body"] + 3, "color": c["text_light"], "bold": True},
            "role":            {"t": "text", "text": "Rol Profesional", "pt": f["size_body"], "color": c["accent"]},
            "seniority":       {"t": "text", "text": "Senior  |  8 anos de exp.", "pt": f["size_caption"] + 1, "color": c["text_muted"]},
            "panel_divider":   {"t": "fill", "fill": c["accent"]},
            "skills_label":    {"t": "text", "text": "HABILIDADES TECNICAS", "pt": f["size_caption"], "color": c["accent"], "bold": True},
            "skills_area":     {"t": "fill", "fill": c["secondary"], "label": "[ TAGS DE HABILIDADES ]"},
            "highlights_label": {"t": "text", "text": "LOGROS DESTACADOS", "pt": f["size_caption"], "color": c["accent"], "bold": True},
            "highlights_area":  {"t": "fill", "fill": c["divider"], "label": "[ LISTA DE LOGROS ]"},
        },
        "stat_callout": {
            "_bg": c["background"],
            "accent_bar_top":    {"t": "fill", "fill": c["accent"]},
            "accent_bar_bottom": {"t": "fill", "fill": c["accent"]},
            "big_number":        {"t": "text", "text": "98%", "pt": f["size_stat"], "color": c["accent"], "bold": True, "align": "center"},
            "label":             {"t": "text", "text": "Descripcion del indicador clave", "pt": f["size_subheading"], "color": c["text_dark"], "bold": True, "align": "center"},
            "divider":           {"t": "fill", "fill": c["accent"]},
            "context":           {"t": "text", "text": "Contexto adicional del dato estadistico", "pt": f["size_body"], "color": c["text_muted"], "align": "center"},
            "source":            {"t": "text", "text": "Fuente: Area de Datos", "pt": f["size_caption"], "color": c["text_muted"], "align": "right", "italic": True},
        },
        "closing": {
            "_bg": c["background"],
            "accent_bar_top":    {"t": "fill", "fill": c["accent"]},
            "accent_bar_bottom": {"t": "fill", "fill": c["accent"]},
            "logo":              {"t": "image", "label": "LOGO"},
            "headline":          {"t": "text", "text": "Listo para transformar tu operacion?", "pt": f["size_heading"], "color": c["text_dark"], "bold": True, "align": "center"},
            "divider":           {"t": "fill", "fill": c["accent"]},
            "contact_label":     {"t": "text", "text": "CONTACTO", "pt": f["size_caption"], "color": c["accent"], "bold": True, "align": "center"},
            "contact_name":      {"t": "text", "text": "Nombre del Contacto", "pt": f["size_body"], "color": c["text_dark"], "bold": True, "align": "center"},
            "contact_email":     {"t": "text", "text": "email@softgic.com", "pt": f["size_body"], "color": c["accent"], "align": "center"},
            "contact_phone":     {"t": "text", "text": "+57 300 000 0000", "pt": f["size_body"], "color": c["text_dark"], "align": "center"},
            "cta_box":           {"t": "fill", "fill": c["accent"]},
            "cta_text":          {"t": "text", "text": "Llamado a la accion del cierre", "pt": f["size_body"], "color": c["primary"], "bold": True, "align": "center"},
        },
        "grafico": {
            "_bg": c["background"],
            "header_bar":  {"t": "fill", "fill": c["primary"]},
            "title":       {"t": "text", "text": "Titulo del Grafico", "pt": f["size_subheading"], "color": c["text_light"], "bold": True},
            "accent_line": {"t": "fill", "fill": c["accent"]},
            "chart_area":  {"t": "fill", "fill": c["divider"], "label": "AREA DEL GRAFICO  (chart_type: column / bar / line / pie)"},
            "source":      {"t": "text", "text": "Fuente: Nombre de la fuente", "pt": f["size_caption"], "color": c["text_muted"], "italic": True},
        },
    }


_SLIDE_ORDER = [
    "cover", "section_divider", "content_two_col", "content_one_col",
    "pricing_table", "profile_card", "stat_callout", "closing", "grafico",
]


def _add_label_bar(slide, slide_type: str) -> None:
    """Barra negra en la parte superior con instrucciones."""
    shape = slide.shapes.add_shape(_RECT, Inches(0), Inches(0), Inches(W), Inches(0.32))
    shape.name = "_LABEL_"
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor(0, 0, 0)
    shape.line.fill.background()
    tf = shape.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = f"LAYOUT: {slide_type}   |   Mueve, redimensiona o cambia colores y fuentes. Guarda (Ctrl+S) y cierra."
    run.font.size = Pt(8)
    run.font.bold = True
    run.font.color.rgb = RGBColor(255, 220, 50)


def _apply_shape_style(shape, name: str, fill_hex: str) -> None:
    """Aplica nombre, relleno solido y sin borde a un shape ya creado."""
    shape.name = name
    shape.fill.solid()
    shape.fill.fore_color.rgb = _rgb(fill_hex)
    shape.line.fill.background()


def _draw_custom_shape(slide, el_name: str, stored: dict, fill: str) -> bool:
    """Dibuja un shape personalizado (kind: shape) usando shape_type_id o shape_prst.

    Retorna True si el shape se creo correctamente, False si hay que usar fallback.
    Soporta formas cuyo prst no esta en el enum de python-pptx (Office 2019+,
    snip corners, etc.) mediante parcheo directo del XML del shape.
    """
    left   = stored["left"]
    top    = stored["top"]
    width  = stored["width"]
    height = stored["height"]

    # Intento 1: MSO_AUTO_SHAPE_TYPE entero (reconocido por python-pptx)
    shape_type_id = stored.get("shape_type_id")
    if isinstance(shape_type_id, int):
        try:
            shape = slide.shapes.add_shape(
                shape_type_id,
                Inches(left), Inches(top), Inches(width), Inches(height),
            )
            _apply_shape_style(shape, el_name, fill)
            return True
        except Exception:
            pass

    # Intento 2: prst string — crear rect y parchear XML
    shape_prst = stored.get("shape_prst")
    if shape_prst:
        try:
            from pptx.oxml.ns import qn
            shape = slide.shapes.add_shape(
                _RECT,
                Inches(left), Inches(top), Inches(width), Inches(height),
            )
            prstGeom = shape._element.spPr.find(qn('a:prstGeom'))
            if prstGeom is not None:
                prstGeom.set('prst', shape_prst)
                avLst = prstGeom.find(qn('a:avLst'))
                if avLst is not None:
                    avLst.clear()
            _apply_shape_style(shape, el_name, fill)
            return True
        except Exception:
            pass

    return False


def _add_fill_element(slide, name: str, pos: dict, el_def: dict) -> None:
    """Rectangulo de color (barra, panel, linea divisora)."""
    shape = slide.shapes.add_shape(
        _RECT,
        Inches(pos["left"]), Inches(pos["top"]),
        Inches(pos["width"]), Inches(pos["height"]),
    )
    shape.name = name

    fill_hex = el_def.get("fill", "CCCCCC").lstrip("#")
    shape.fill.solid()
    shape.fill.fore_color.rgb = _rgb(fill_hex)
    shape.line.fill.background()

    tf = shape.text_frame
    tf.word_wrap = True
    tf.clear()
    label = el_def.get("label", name)
    if pos["height"] >= 0.25:
        para = tf.paragraphs[0]
        para.alignment = PP_ALIGN.CENTER
        run = para.add_run()
        run.text = label
        run.font.size = Pt(7)
        run.font.bold = True
        run.font.color.rgb = _label_color(fill_hex)


def _add_text_element(slide, name: str, pos: dict, el_def: dict, font_family: str) -> None:
    """Caja de texto con contenido de muestra en fuente/tamano/color reales.
    Si el_def contiene 'runs', renderiza texto multicolor respetando color por run."""
    txBox = slide.shapes.add_textbox(
        Inches(pos["left"]), Inches(pos["top"]),
        Inches(pos["width"]), Inches(pos["height"]),
    )
    txBox.name = name

    align_map = {
        "left":   PP_ALIGN.LEFT,
        "center": PP_ALIGN.CENTER,
        "right":  PP_ALIGN.RIGHT,
    }
    align      = align_map.get(el_def.get("align", "left"), PP_ALIGN.LEFT)
    default_color = str(el_def.get("color", "333333")).lstrip("#")
    font_pt    = el_def.get("pt", 12)
    bold       = el_def.get("bold", False)
    italic     = el_def.get("italic", False)
    runs_data  = el_def.get("runs")

    tf = txBox.text_frame
    tf.word_wrap = True
    tf.clear()

    if runs_data:
        # Texto multicolor: un run por entrada con su propio color
        para = tf.paragraphs[0]
        para.alignment = align
        for run_info in runs_data:
            run = para.add_run()
            run.text = str(run_info.get("text", ""))
            run.font.name = font_family
            run.font.size = Pt(int(run_info.get("font_size", font_pt)))
            run.font.bold = bool(run_info.get("bold", bold))
            run.font.italic = italic
            run_color = str(run_info.get("color", default_color)).lstrip("#")
            try:
                run.font.color.rgb = _rgb(run_color)
            except Exception:
                run.font.color.rgb = _rgb(default_color)
    else:
        # Texto de un solo color (caso normal)
        sample_text = el_def.get("text", name)
        lines = sample_text.split("\n")
        for i, line in enumerate(lines):
            para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            para.alignment = align
            run = para.add_run()
            run.text = line
            run.font.name = font_family
            run.font.size = Pt(font_pt)
            run.font.bold = bold
            run.font.italic = italic
            try:
                run.font.color.rgb = _rgb(default_color)
            except Exception:
                run.font.color.rgb = _rgb("333333")


def _apply_gradient_bg(slide, angle: int, stops: list) -> None:
    """Aplica fondo degradado lineal al slide via XML."""
    import lxml.etree as _etree
    from pptx.oxml.ns import qn

    slide.background.fill.solid()
    bg_pr = slide._element.find('.//' + qn('p:bgPr'))
    if bg_pr is None:
        return

    for tag in (qn('a:solidFill'), qn('a:gradFill'), qn('a:noFill'), qn('a:pattFill'), qn('a:blipFill')):
        el = bg_pr.find(tag)
        if el is not None:
            bg_pr.remove(el)

    grad_fill = _etree.Element(qn('a:gradFill'))
    gs_lst = _etree.SubElement(grad_fill, qn('a:gsLst'))
    for stop in stops:
        pos_pct = int(stop.get("pos", 0))
        color   = str(stop.get("color", "000000")).upper().lstrip("#")
        gs = _etree.SubElement(gs_lst, qn('a:gs'))
        gs.set("pos", str(pos_pct * 1000))
        srgb = _etree.SubElement(gs, qn('a:srgbClr'))
        srgb.set("val", color)

    lin = _etree.SubElement(grad_fill, qn('a:lin'))
    lin.set("ang", str(angle * 60000))
    lin.set("scaled", "0")
    bg_pr.insert(0, grad_fill)


def _set_slide_bg(slide, bg_hex_default: str, slide_layout: dict) -> None:
    """Aplica el fondo del slide. Usa _background de layout.yaml si existe (cambio
    intencional del usuario via Format Background), sino usa el color del tema."""
    bg_override = slide_layout.get("_background")

    if isinstance(bg_override, str) and len(bg_override) == 6:
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = _rgb(bg_override)
    elif isinstance(bg_override, dict):
        bg_type = bg_override.get("type", "solid")
        if bg_type == "solid":
            color = str(bg_override.get("color", bg_hex_default)).lstrip("#")
            slide.background.fill.solid()
            slide.background.fill.fore_color.rgb = _rgb(color)
        elif bg_type == "gradient":
            _apply_gradient_bg(slide, int(bg_override.get("angle", 90)), bg_override.get("stops", []))
    else:
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = _rgb(bg_hex_default)


def _add_image_placeholder(slide, name: str, pos: dict, el_def: dict) -> None:
    """Rectangulo gris para logotipos o fotos."""
    shape = slide.shapes.add_shape(
        _RECT,
        Inches(pos["left"]), Inches(pos["top"]),
        Inches(pos["width"]), Inches(pos["height"]),
    )
    shape.name = name
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor(170, 170, 170)
    shape.line.fill.background()

    label = el_def.get("label", name.upper())
    tf = shape.text_frame
    tf.clear()
    para = tf.paragraphs[0]
    para.alignment = PP_ALIGN.CENTER
    run = para.add_run()
    run.text = label
    run.font.size = Pt(8)
    run.font.bold = True
    run.font.color.rgb = RGBColor(255, 255, 255)


def _render_slide(prs: Presentation, slide_type: str, layout_data: dict, defs: dict,
                  font_family: str, default_positions: dict) -> None:
    blank = prs.slide_layouts[6]
    slide = prs.slides.add_slide(blank)

    slide_def = defs.get(slide_type, {})
    bg_color = slide_def.get("_bg", "FFFFFF").lstrip("#")

    slide_layout = layout_data.get(slide_type, {})
    _set_slide_bg(slide, bg_color, slide_layout)

    _add_label_bar(slide, slide_type)

    # Mostrar elementos canonicos del renderer (_D), respetando el estado hidden.
    # Si el usuario borro un elemento (hidden:true), NO se incluye en el PPTX — asi
    # al volver a correr apply sigue marcado como eliminado. Para restaurarlo, el
    # usuario agrega una forma en PowerPoint con el nombre exacto del elemento.
    for el_name, d_pos in default_positions.items():
        stored = slide_layout.get(el_name, {})

        # Respetar eliminaciones: si el usuario borro este elemento, no lo mostrar
        if isinstance(stored, dict) and stored.get("hidden"):
            continue

        # Usar posicion de layout.yaml si es valida, sino el default de _D
        if isinstance(stored, dict) and all(k in stored for k in ("left", "top", "width", "height")):
            pos = stored
        else:
            pos = d_pos

        el_def = slide_def.get(el_name, {"t": "fill", "fill": "CCCCCC"})

        # Aplicar color/tamano/negrita/relleno almacenados en layout.yaml al el_def
        # para que el editor visual refleje los cambios del usuario.
        if isinstance(stored, dict) and not stored.get("hidden"):
            overrides: dict = {}
            if "color" in stored:
                overrides["color"] = stored["color"]
            if "font_size" in stored:
                overrides["pt"] = stored["font_size"]
            if stored.get("bold") is not None:
                overrides["bold"] = stored["bold"]
            if "fill" in stored:
                overrides["fill"] = stored["fill"]
            if overrides:
                el_def = {**el_def, **overrides}

        el_type = el_def.get("t", "fill")

        if el_type == "text":
            _add_text_element(slide, el_name, pos, el_def, font_family)
        elif el_type == "image":
            _add_image_placeholder(slide, el_name, pos, el_def)
        else:
            _add_fill_element(slide, el_name, pos, el_def)

    # Mostrar formas extra (no en _D) que el usuario haya agregado y guardado
    for el_name, stored in slide_layout.items():
        if el_name in default_positions:
            continue  # ya mostrado arriba
        if el_name == "background_color":
            continue  # elemento legacy del editor visual anterior
        if not isinstance(stored, dict) or stored.get("hidden"):
            continue
        if not all(k in stored for k in ("left", "top", "width", "height")):
            continue

        kind = stored.get("kind", "box")
        fill = stored.get("fill", "AAAAAA")

        # Fallback por nombre: ovaIos guardados sin campo 'kind' en versiones anteriores
        if kind == "box" and any(h in el_name.lower() for h in ("oval", "ellipse", "circle", "circulo")):
            kind = "oval"

        if kind == "image":
            src = stored.get("src", "")
            img_path = (ROOT / src) if src and not Path(src).is_absolute() else Path(src) if src else None
            if img_path and img_path.exists():
                try:
                    pic = slide.shapes.add_picture(
                        str(img_path),
                        Inches(stored["left"]), Inches(stored["top"]),
                        Inches(stored["width"]), Inches(stored["height"]),
                    )
                    pic.name = el_name
                except Exception:
                    _add_fill_element(slide, el_name, stored, {"t": "fill", "fill": "999999", "label": f"IMG: {el_name}"})
            else:
                _add_fill_element(slide, el_name, stored, {"t": "fill", "fill": "BBBBBB", "label": f"IMG no encontrada: {el_name}"})
        elif kind == "oval":
            shape = slide.shapes.add_shape(
                _OVAL,
                Inches(stored["left"]), Inches(stored["top"]),
                Inches(stored["width"]), Inches(stored["height"]),
            )
            _apply_shape_style(shape, el_name, fill)
        elif kind == "text":
            try:
                text_def = {
                    "text":  stored.get("text", el_name),
                    "pt":    stored.get("font_size", 12),
                    "color": stored.get("color", "333333"),
                    "bold":  stored.get("bold", False),
                    "runs":  stored.get("runs"),
                }
                _add_text_element(slide, el_name, stored, text_def, font_family)
            except Exception as _e:
                _add_fill_element(slide, el_name, stored, {"t": "fill", "fill": "CCCCCC", "label": el_name})
        elif kind == "shape":
            try:
                if not _draw_custom_shape(slide, el_name, stored, fill):
                    _add_fill_element(slide, el_name, stored, {"t": "fill", "fill": fill, "label": el_name})
            except Exception:
                _add_fill_element(slide, el_name, stored, {"t": "fill", "fill": fill, "label": el_name})
        else:
            try:
                _add_fill_element(slide, el_name, stored, {"t": "fill", "fill": fill, "label": el_name})
            except Exception:
                pass


def _load_slide_defaults() -> dict[str, dict]:
    """Carga los dicts _D de cada renderer para saber que elementos mostrar en el editor."""
    if str(ROOT) not in sys.path:
        sys.path.insert(0, str(ROOT))
    try:
        from src.layout.renderers.cover          import _D as _D_cover
        from src.layout.renderers.section_divider import _D as _D_section_divider
        from src.layout.renderers.content_two_col import _D as _D_content_two_col
        from src.layout.renderers.content_one_col import _D as _D_content_one_col
        from src.layout.renderers.pricing_table   import _D as _D_pricing_table
        from src.layout.renderers.profile_card    import _D as _D_profile_card
        from src.layout.renderers.stat_callout    import _D as _D_stat_callout
        from src.layout.renderers.closing         import _D as _D_closing
        from src.layout.renderers.grafico         import _D as _D_grafico
        return {
            "cover":            _D_cover,
            "section_divider":  _D_section_divider,
            "content_two_col":  _D_content_two_col,
            "content_one_col":  _D_content_one_col,
            "pricing_table":    _D_pricing_table,
            "profile_card":     _D_profile_card,
            "stat_callout":     _D_stat_callout,
            "closing":          _D_closing,
            "grafico":          _D_grafico,
        }
    except ImportError as e:
        print(f"AVISO: no se pudo importar _D de los renderers ({e})")
        return {}


def main() -> None:
    layout_data: dict = {}
    if LAYOUT_YAML.exists():
        layout_data = yaml.safe_load(LAYOUT_YAML.read_text(encoding="utf-8")) or {}

    theme = _load_theme()
    c_raw = theme.get("colors", {})
    f_raw = theme.get("fonts", {})

    c = {
        "primary":    c_raw.get("primary",    "#0A1628").lstrip("#"),
        "secondary":  c_raw.get("secondary",  "#1E3A5F").lstrip("#"),
        "accent":     c_raw.get("accent",     "#00AEEF").lstrip("#"),
        "background": c_raw.get("background", "#FFFFFF").lstrip("#"),
        "text_dark":  c_raw.get("text_dark",  "#0A1628").lstrip("#"),
        "text_light": c_raw.get("text_light", "#FFFFFF").lstrip("#"),
        "text_muted": c_raw.get("text_muted", "#6B7280").lstrip("#"),
        "divider":    c_raw.get("divider",    "#E5E7EB").lstrip("#"),
    }
    f = {
        "family":          f_raw.get("family",         "Calibri"),
        "size_heading":    f_raw.get("size_heading",   36),
        "size_subheading": f_raw.get("size_subheading",24),
        "size_body":       f_raw.get("size_body",      14),
        "size_caption":    f_raw.get("size_caption",   10),
        "size_stat":       f_raw.get("size_stat",      72),
    }

    defs = _build_defs(c, f)
    slide_defaults = _load_slide_defaults()

    prs = Presentation()
    prs.slide_width  = Inches(W)
    prs.slide_height = Inches(H)

    for slide_type in _SLIDE_ORDER:
        default_positions = slide_defaults.get(slide_type, {})
        _render_slide(prs, slide_type, layout_data, defs, f["family"], default_positions)
        print(f"  OK {slide_type}")

    OUT_PPTX.parent.mkdir(exist_ok=True)
    prs.save(str(OUT_PPTX))

    print()
    print(f"Editor visual guardado en: {OUT_PPTX}")
    print("(archivo en config/ — no es un asset del deck, es una herramienta de diseno)")
    print()
    print("Instrucciones:")
    print("  - Mueve y redimensiona los elementos donde quieras")
    print("  - Cambia el color de relleno de las barras y paneles")
    print("  - Cambia el tamano de fuente del texto de ejemplo")
    print("  - Cambia el color del texto de ejemplo")
    print("  - NO cambies el nombre de los elementos (visible en 'Selection Pane')")
    print("  - Guarda con Ctrl+S y cierra PowerPoint")
    print("  - Luego usa 'Aplicar cambios del editor visual' para actualizar el proyecto")

    if os.name == "nt":
        os.startfile(str(OUT_PPTX))  # type: ignore[attr-defined]


if __name__ == "__main__":
    main()
