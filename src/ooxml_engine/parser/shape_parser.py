"""shape (python-pptx) -> OoxmlElement. Alcance M1: PICTURE, AUTO_SHAPE,
TEXT_BOX y PLACEHOLDER se parsean en detalle. Cualquier otro tipo (GROUP,
TABLE nativa, CHART nativo, CONNECTOR, MEDIA, etc.) se modela como
OoxmlElement base con kind=UNKNOWN — sin perder posicion/nombre, pero sin
contenido detallado, con un log de aviso (nunca una excepcion). Se agrega
soporte detallado para esos tipos en M2 si el template que se este
procesando realmente los usa (config/template_master.pptx no los usa).
"""
from __future__ import annotations

import logging

from pptx.enum.shapes import MSO_SHAPE_TYPE

from src.ooxml_engine.model.elements import (
    ElementKind,
    ImageElement,
    OoxmlElement,
    ShapeElement,
    TextElement,
    Transform,
)
from src.ooxml_engine.model.style import Fill, FillKind, Paragraph
from src.ooxml_engine.parser.style_parser import parse_fill
from src.ooxml_engine.parser.text_parser import parse_text_frame
from src.ooxml_engine.parser.xml_readers import get_prst_geometry, read_text_margins_emu

logger = logging.getLogger("softgic.ooxml_engine.shape_parser")


def _build_transform(shape) -> Transform:
    rotation = 0.0
    try:
        rotation = round(float(shape.rotation), 2)
    except (AttributeError, ValueError):
        pass
    return Transform(
        left_emu=int(shape.left or 0),
        top_emu=int(shape.top or 0),
        width_emu=int(shape.width or 0),
        height_emu=int(shape.height or 0),
        rotation_deg=rotation,
    )


def _placeholder_info(shape) -> tuple[bool, str | None]:
    try:
        if not shape.is_placeholder:
            return False, None
        return True, shape.placeholder_format.type.name
    except (AttributeError, ValueError, KeyError):
        return bool(getattr(shape, "is_placeholder", False)), None


def _build_text_element(
    shape, transform: Transform, is_ph: bool, ph_type: str | None,
    paragraphs: list[Paragraph] | None = None, fill: Fill | None = None,
) -> TextElement:
    margins = read_text_margins_emu(shape)
    word_wrap = True
    try:
        word_wrap = bool(shape.text_frame.word_wrap) if shape.text_frame.word_wrap is not None else True
    except AttributeError:
        pass
    return TextElement(
        shape_id=shape.shape_id,
        name=shape.name,
        kind=ElementKind.TEXT,
        transform=transform,
        is_placeholder=is_ph,
        placeholder_type=ph_type,
        paragraphs=paragraphs if paragraphs is not None else parse_text_frame(shape),
        fill=fill if fill is not None else parse_fill(shape),
        margin_left_emu=margins[0],
        margin_top_emu=margins[1],
        margin_right_emu=margins[2],
        margin_bottom_emu=margins[3],
        word_wrap=word_wrap,
    )


def parse_shape(shape) -> OoxmlElement:
    transform = _build_transform(shape)
    is_ph, ph_type = _placeholder_info(shape)

    try:
        shape_type = shape.shape_type
    except (AttributeError, ValueError, NotImplementedError):
        shape_type = None

    if shape_type == MSO_SHAPE_TYPE.PICTURE:
        try:
            image = shape.image
            return ImageElement(
                shape_id=shape.shape_id,
                name=shape.name,
                kind=ElementKind.IMAGE,
                transform=transform,
                is_placeholder=is_ph,
                placeholder_type=ph_type,
                blob=image.blob,
                image_format=image.ext or "png",
                crop_left=float(shape.crop_left or 0.0),
                crop_right=float(shape.crop_right or 0.0),
                crop_top=float(shape.crop_top or 0.0),
                crop_bottom=float(shape.crop_bottom or 0.0),
            )
        except (AttributeError, ValueError) as exc:
            logger.warning("No se pudo leer imagen '%s': %s", shape.name, exc)
            return OoxmlElement(
                shape_id=shape.shape_id, name=shape.name, kind=ElementKind.UNKNOWN,
                transform=transform, is_placeholder=is_ph, placeholder_type=ph_type,
            )

    if shape_type == MSO_SHAPE_TYPE.TEXT_BOX:
        return _build_text_element(shape, transform, is_ph, ph_type)

    if shape_type in (MSO_SHAPE_TYPE.AUTO_SHAPE, MSO_SHAPE_TYPE.PLACEHOLDER):
        fill = parse_fill(shape)
        text_paragraphs = parse_text_frame(shape) if getattr(shape, "has_text_frame", False) else []
        has_real_text = any(p.runs for p in text_paragraphs)

        if fill.kind == FillKind.NONE and has_real_text:
            return _build_text_element(shape, transform, is_ph, ph_type, text_paragraphs, fill)

        text_sub: TextElement | None = None
        if has_real_text:
            text_sub = _build_text_element(shape, transform, is_ph, ph_type, text_paragraphs, fill)

        return ShapeElement(
            shape_id=shape.shape_id,
            name=shape.name,
            kind=ElementKind.SHAPE,
            transform=transform,
            is_placeholder=is_ph,
            placeholder_type=ph_type,
            geometry_preset=get_prst_geometry(shape) or "rect",
            fill=fill,
            text=text_sub,
        )

    logger.warning(
        "Shape '%s' de tipo %s no soportado en M1 — modelado como UNKNOWN", shape.name, shape_type
    )
    return OoxmlElement(
        shape_id=shape.shape_id, name=shape.name, kind=ElementKind.UNKNOWN,
        transform=transform, is_placeholder=is_ph, placeholder_type=ph_type,
    )
