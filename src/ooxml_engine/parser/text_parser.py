"""Parseo de un text_frame a list[Paragraph]. Generaliza la logica de
extraccion de runs/parrafos de apply_visual_config.py::_extract_text_content
pero SIN la heuristica "primer valor encontrado a nivel shape" — aqui cada
run conserva su propio FontSpec siempre (no solo cuando hay multicolor),
porque el modelo no necesita colapsar a un unico font_size/color por shape.
"""
from __future__ import annotations

from pptx.oxml.ns import qn

from src.ooxml_engine.model.style import FontSpec, Paragraph, TextRun
from src.ooxml_engine.parser.xml_readers import read_autofit_font_scale, resolve_scheme_color

_ALIGN_MAP = {
    1: "left",  # PP_ALIGN.LEFT
    2: "center",
    3: "right",
    4: "justify",
}


def _read_run_size_pt(run, autofit_scale: float) -> float | None:
    try:
        sz_emu = run.font.size
        if sz_emu is not None:
            pt = int(sz_emu) / 12700
            return round(pt * autofit_scale, 1) if pt > 0 else None
    except (AttributeError, ValueError):
        pass
    try:
        rpr = run._r.find(qn("a:rPr"))
        if rpr is not None:
            sz_str = rpr.get("sz")
            if sz_str:
                pt = int(sz_str) / 100
                return round(pt * autofit_scale, 1)
    except (AttributeError, ValueError):
        pass
    return None


def _read_run_color_hex(run, shape) -> str | None:
    try:
        clr = run.font.color
        if clr and clr.type is not None:
            s = str(clr.rgb).lstrip("#").upper()
            return s if len(s) == 6 else None
    except (AttributeError, ValueError, KeyError, TypeError):
        pass
    try:
        rpr = run._r.find(qn("a:rPr"))
        if rpr is None:
            return None
        sf = rpr.find(qn("a:solidFill"))
        if sf is None:
            return None
        srgb = sf.find(qn("a:srgbClr"))
        if srgb is not None:
            v = srgb.get("val", "")
            return v.upper() if len(v) == 6 else None
        sc = sf.find(qn("a:schemeClr"))
        if sc is not None:
            lm_e = sc.find(qn("a:lumMod"))
            lo_e = sc.find(qn("a:lumOff"))
            lm = int(lm_e.get("val", "100000")) if lm_e is not None else 100000
            lo = int(lo_e.get("val", "0")) if lo_e is not None else 0
            return resolve_scheme_color(shape, sc.get("val", ""), lm, lo)
    except (AttributeError, ValueError):
        pass
    return None


def _read_bullet_char(para) -> str | None:
    try:
        ppr = para._p.find(qn("a:pPr"))
        if ppr is None:
            return None
        bu_char = ppr.find(qn("a:buChar"))
        if bu_char is not None:
            return bu_char.get("char")
    except AttributeError:
        pass
    return None


def parse_text_frame(shape) -> list[Paragraph]:
    """shape: cualquier shape con has_text_frame == True. Retorna [] si no
    hay texto real (texto vacio/solo espacios) — un text_frame tecnicamente
    presente pero vacio no produce parrafos fantasma.
    """
    try:
        if not shape.has_text_frame:
            return []
    except AttributeError:
        return []

    if not shape.text_frame.text.strip():
        return []

    autofit_scale = read_autofit_font_scale(shape)
    paragraphs: list[Paragraph] = []

    for para in shape.text_frame.paragraphs:
        runs: list[TextRun] = []
        for run in para.runs:
            if not run.text:
                continue
            font = FontSpec(
                size_pt=_read_run_size_pt(run, autofit_scale),
                bold=bool(run.font.bold),
                italic=bool(run.font.italic),
                underline=bool(run.font.underline),
                color_hex=_read_run_color_hex(run, shape),
            )
            try:
                font.family = run.font.name
            except AttributeError:
                pass
            runs.append(TextRun(text=run.text, font=font))

        alignment = _ALIGN_MAP.get(int(para.alignment) if para.alignment is not None else 1, "left")
        paragraphs.append(
            Paragraph(runs=runs, alignment=alignment, bullet_char=_read_bullet_char(para))
        )

    return paragraphs
