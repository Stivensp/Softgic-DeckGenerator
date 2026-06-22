"""Lectores XML crudo reutilizables para propiedades que python-pptx 1.0.2
no expone con setter/getter publico. Generalizacion de las funciones
equivalentes en tools/apply_visual_config.py (_resolve_scheme_color,
_read_autofit_scale, _get_shape_prst, _read_text_margins) — aqui NO estan
atadas a los 9 tipos de slide fijos, reciben cualquier shape/elemento.

Convencion: cada funcion retorna None (o un valor por defecto explicito,
documentado) si la propiedad no existe en el XML — nunca lanza, salvo que
el XML este corrupto de forma que ni python-pptx pudo abrirlo (eso ya
habria fallado antes, en Presentation(path)).

Leccion ya aprendida y corregida esta sesion en apply_visual_config.py:
shape.auto_shape_type lanza ValueError (NO devuelve None) cuando el prst
no esta en el enum de python-pptx — cualquier funcion de aqui que toque
esa propiedad debe capturar ValueError explicitamente, no solo
AttributeError/TypeError/KeyError.
"""
from __future__ import annotations

import colorsys

from pptx.oxml.ns import qn

EMU_PER_INCH = 914400

_RT = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/"
_SCHEME_ALIASES = {"tx1": "dk1", "tx2": "dk2", "bg1": "lt1", "bg2": "lt2"}


def get_prst_geometry(shape) -> str | None:
    """Lee el 'prst' (preset geometry) directamente del XML del shape.

    Usado como fallback cuando auto_shape_type no esta en el enum de
    python-pptx (formas de Office 2019+, snip corners, etc.) — y tambien
    como via principal aqui, ya que es mas simple leer el string crudo una
    vez que confiar en que el enum lo reconozca.
    """
    try:
        prst_geom = shape._element.spPr.find(qn("a:prstGeom"))
        return prst_geom.get("prst") if prst_geom is not None else None
    except (AttributeError, ValueError):
        return None


def resolve_scheme_color(
    part_or_shape, scheme_name: str, lum_mod: int = 100000, lum_off: int = 0
) -> str | None:
    """Convierte un schemeClr de PowerPoint a hex RGB navegando
    slide->layout->master->tema. Aplica luminancia modificada (lumMod,
    lumOff) para cubrir las variantes mas claro/mas oscuro del selector
    de colores de PowerPoint.
    """
    try:
        import lxml.etree as _etree

        elem = _SCHEME_ALIASES.get(scheme_name, scheme_name)
        if elem == "phClr":
            return None

        part = getattr(part_or_shape, "part", part_or_shape)
        layout = part.part_related_by(_RT + "slideLayout")
        master = layout.part_related_by(_RT + "slideMaster")
        theme = master.part_related_by(_RT + "theme")

        theme_xml = _etree.fromstring(theme.blob)
        scheme = theme_xml.find(".//" + qn("a:clrScheme"))
        if scheme is None:
            return None

        clr = scheme.find(qn("a:" + elem))
        if clr is None:
            return None

        srgb = clr.find(qn("a:srgbClr"))
        base = srgb.get("val", "") if srgb is not None else ""
        if not base:
            sys_c = clr.find(qn("a:sysClr"))
            base = sys_c.get("lastClr", "") if sys_c is not None else ""
        if len(base) != 6:
            return None

        if lum_mod == 100000 and lum_off == 0:
            return base.upper()

        r, g, b = int(base[0:2], 16) / 255, int(base[2:4], 16) / 255, int(base[4:6], 16) / 255
        h, l, s = colorsys.rgb_to_hls(r, g, b)
        l = max(0.0, min(1.0, l * (lum_mod / 100000) + (lum_off / 100000)))
        r2, g2, b2 = colorsys.hls_to_rgb(h, l, s)
        return f"{round(r2 * 255):02X}{round(g2 * 255):02X}{round(b2 * 255):02X}"
    except (AttributeError, KeyError, ValueError):
        return None


def read_solid_fill_from_parent(parent_elem, theme_anchor) -> str | None:
    """Busca un <a:solidFill> hijo directo de parent_elem (spPr de un shape,
    o bgPr de un slide) y resuelve su color a hex — sRGB explicito o
    schemeClr (color de tema, resuelto via theme_anchor: el shape o el
    slide del que colgar part_related_by para llegar al theme part).
    """
    try:
        if parent_elem is None:
            return None
        sf = parent_elem.find(qn("a:solidFill"))
        if sf is None:
            return None

        srgb = sf.find(qn("a:srgbClr"))
        if srgb is not None:
            v = srgb.get("val", "")
            return v.upper() if len(v) == 6 else None

        sc = sf.find(qn("a:schemeClr"))
        if sc is not None:
            lm_e = sc.find(qn("a:lumMod"))
            lo_e = sc.find(qn("a:lumOff"))
            lm = int(lm_e.get("val", "100000")) if lm_e is not None else 100000
            lo = int(lo_e.get("val", "0")) if lo_e is not None else 0
            return resolve_scheme_color(theme_anchor, sc.get("val", ""), lm, lo)
    except (AttributeError, ValueError):
        pass
    return None


def read_fill_hex(shape) -> str | None:
    """Lee el color de relleno solido del shape como hex RGB. Maneja tanto
    sRGB explicito como schemeClr (colores de tema)."""
    try:
        return read_solid_fill_from_parent(shape._element.spPr, shape)
    except AttributeError:
        return None


def read_autofit_font_scale(shape) -> float:
    """Lee el factor de escala que PowerPoint aplica via 'Reducir texto al
    desbordar' (normAutofit/@fontScale). El tamano de fuente en el XML del
    run es el NOMINAL configurado por el usuario — si el texto no entraba
    en la forma, PowerPoint lo encoge visualmente sin tocar ese valor
    nominal. Sin aplicar este factor, el texto se regenera mas grande de
    lo que realmente se veia.
    """
    try:
        body_pr = shape.text_frame._txBody.find(qn("a:bodyPr"))
        if body_pr is None:
            return 1.0
        norm_autofit = body_pr.find(qn("a:normAutofit"))
        if norm_autofit is None:
            return 1.0
        scale = norm_autofit.get("fontScale")
        return int(scale) / 100000 if scale else 1.0
    except (AttributeError, ValueError):
        return 1.0


def read_text_margins_emu(shape) -> tuple[int, int, int, int]:
    """Lee los margenes internos (insets) del text_frame en EMU crudo (no
    pulgadas — la conversion es responsabilidad de quien consuma el
    modelo). Retorna (left, top, right, bottom). python-pptx ya devuelve
    los defaults documentados de PowerPoint (91440/45720 EMU) cuando no
    estan definidos explicitamente en el XML, asi que esto nunca falla en
    la practica, pero se documenta el fallback (0,0,0,0) por completitud.
    """
    try:
        tf = shape.text_frame
        return (
            int(tf.margin_left or 0),
            int(tf.margin_top or 0),
            int(tf.margin_right or 0),
            int(tf.margin_bottom or 0),
        )
    except AttributeError:
        return (0, 0, 0, 0)
