"""Entrypoint del parser: parse_pptx(path) -> OoxmlPresentation."""
from __future__ import annotations

from pathlib import Path

from pptx import Presentation

from src.ooxml_engine.model.document import OoxmlPresentation
from src.ooxml_engine.parser.slide_parser import parse_slide
from src.ooxml_engine.parser.theme_parser import parse_theme


def parse_pptx(path: str | Path) -> OoxmlPresentation:
    prs = Presentation(str(path))
    theme = parse_theme(prs.slide_masters[0])

    return OoxmlPresentation(
        slide_width_emu=int(prs.slide_width),
        slide_height_emu=int(prs.slide_height),
        theme=theme,
        slides=[parse_slide(slide, i) for i, slide in enumerate(prs.slides)],
        source_path=str(path),
    )
