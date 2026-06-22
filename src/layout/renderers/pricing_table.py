from __future__ import annotations

from typing import Any

from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

from src.layout import layout_config as lc
from src.layout.base_layout import BaseLayoutRenderer
from src.layout.helpers import (
    add_colored_box,
    add_text_box,
    hex_to_rgb,
    render_layout_extras,
    set_slide_background,
    truncate_text,
)
from src.models.slides import PricingTableSlide
from src.models.theme import ThemeModel

_D = {
    "header_bar":  {"left": 0.0, "top": 0.0,  "width": 13.33, "height": 1.1 },
    "title":       {"left": 0.5, "top": 0.15, "width": 12.33, "height": 0.8,  "font_size": 24},
    "accent_line": {"left": 0.0, "top": 1.1,  "width": 13.33, "height": 0.06},
    "table":       {"left": 0.5, "top": 1.3,  "width": 12.33, "height": 0.42},
    "notes":       {"left": 0.5, "top": 0.2,  "width": 12.33, "height": 0.45, "font_size": 10},
}

_COL_WIDTHS = (5.0, 1.1, 1.73, 2.2, 2.3)
_HEADERS = ("Descripcion", "Cant.", "Unidad", "Precio Unit.", "Total")


def _p(el: str) -> dict[str, Any] | None:
    override = lc.pos("pricing_table", el)
    if override is None:
        return None
    return {**_D[el], **override}


class PricingTableRenderer(BaseLayoutRenderer):
    def render(self, slide: object, model: object, theme: ThemeModel) -> None:
        assert isinstance(model, PricingTableSlide)
        c = theme.colors
        f = theme.fonts
        lim = theme.limits

        set_slide_background(slide, c.background)  # type: ignore[arg-type]

        if (p := _p("header_bar")) is not None:
            add_colored_box(slide, p["left"], p["top"], p["width"], p["height"], p.get("fill", c.primary))  # type: ignore[arg-type]

        title = truncate_text(model.title, lim.title_max_chars, "pricing_table.title")
        if (p := _p("title")) is not None:
            add_text_box(slide, p["left"], p["top"], p["width"], p["height"], title,  # type: ignore[arg-type]
                         f.family, p.get("font_size", f.size_subheading), p.get("color", c.text_light),
                         bold=p.get("bold", f.bold_headings))

        if (p := _p("accent_line")) is not None:
            add_colored_box(slide, p["left"], p["top"], p["width"], p["height"], p.get("fill", c.accent))  # type: ignore[arg-type]

        # notes_top: posicion absoluta donde colocar las notas. Si la tabla esta
        # visible, se calcula relativa a su final; si esta oculta, se usa la
        # posicion propia de "notes" sin offset (no hay tabla de referencia).
        notes_top: float | None = None

        pt = _p("table")
        if pt is not None:
            row_h = pt["height"]
            num_rows = 1 + len(model.rows) + 1
            tbl = slide.shapes.add_table(  # type: ignore[union-attr]
                num_rows, 5,
                Inches(pt["left"]), Inches(pt["top"]),
                Inches(pt["width"]), Inches(row_h * num_rows),
            ).table
            for col_i, w in enumerate(_COL_WIDTHS):
                tbl.columns[col_i].width = Inches(w)

            # Header row
            for col_i, header in enumerate(_HEADERS):
                cell = tbl.cell(0, col_i)
                cell.fill.solid()
                cell.fill.fore_color.rgb = hex_to_rgb(c.primary)
                tf = cell.text_frame
                tf.clear()
                para = tf.paragraphs[0]
                para.alignment = PP_ALIGN.CENTER if col_i > 0 else PP_ALIGN.LEFT
                run = para.add_run()
                run.text = header
                run.font.name = f.family
                run.font.size = Pt(f.size_body - 1)
                run.font.bold = True
                run.font.color.rgb = hex_to_rgb(c.text_light)

            # Data rows
            for row_i, row in enumerate(model.rows):
                bg = c.background if row_i % 2 == 0 else c.divider
                values = [
                    truncate_text(row.description, lim.bullet_max_chars, f"rows[{row_i}].description"),
                    str(row.quantity),
                    row.unit or "",
                    f"{model.currency} {row.unit_price:,.2f}",
                    f"{model.currency} {row.total:,.2f}",
                ]
                aligns = [PP_ALIGN.LEFT, PP_ALIGN.CENTER, PP_ALIGN.LEFT, PP_ALIGN.RIGHT, PP_ALIGN.RIGHT]
                for col_i, (val, aln) in enumerate(zip(values, aligns)):
                    cell = tbl.cell(row_i + 1, col_i)
                    cell.fill.solid()
                    cell.fill.fore_color.rgb = hex_to_rgb(bg)
                    tf = cell.text_frame
                    tf.clear()
                    para = tf.paragraphs[0]
                    para.alignment = aln
                    run = para.add_run()
                    run.text = val
                    run.font.name = f.family
                    run.font.size = Pt(f.size_body - 1)
                    run.font.color.rgb = hex_to_rgb(c.text_dark)

            # Totals row
            totals_row = len(model.rows) + 1
            for col_i, (txt, aln) in enumerate(zip(
                ["", "", "", "TOTAL", f"{model.currency} {model.totals:,.2f}"],
                [PP_ALIGN.LEFT, PP_ALIGN.CENTER, PP_ALIGN.LEFT, PP_ALIGN.RIGHT, PP_ALIGN.RIGHT],
            )):
                cell = tbl.cell(totals_row, col_i)
                cell.fill.solid()
                cell.fill.fore_color.rgb = hex_to_rgb(c.accent)
                tf = cell.text_frame
                tf.clear()
                para = tf.paragraphs[0]
                para.alignment = aln
                run = para.add_run()
                run.text = txt
                run.font.name = f.family
                run.font.size = Pt(f.size_body)
                run.font.bold = True
                run.font.color.rgb = hex_to_rgb(c.text_light)

            if (pn := _p("notes")) is not None:
                notes_top = pt["top"] + num_rows * row_h + pn["top"]

        if model.notes and (pn := _p("notes")) is not None:
            top = notes_top if notes_top is not None else pn["top"]
            add_text_box(slide, pn["left"], top, pn["width"], pn["height"],  # type: ignore[arg-type]
                         f"* {model.notes}", f.family, pn.get("font_size", f.size_caption), pn.get("color", c.text_muted), italic=True)

        render_layout_extras(slide, "pricing_table", set(_D.keys()), c.accent, model=model)  # type: ignore[arg-type]
