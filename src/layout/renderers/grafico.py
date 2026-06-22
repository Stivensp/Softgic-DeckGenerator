from __future__ import annotations

from typing import Any

import logging

from pptx.chart.data import ChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches, Pt

from src.layout import layout_config as lc
from src.layout.base_layout import BaseLayoutRenderer
from src.layout.helpers import (
    add_colored_box,
    add_text_box,
    hex_to_rgb,
    render_layout_extras,
    set_slide_background,
    truncate_text,
)
from src.models.slides import GraficoSlide
from src.models.theme import ThemeModel

logger = logging.getLogger("softgic.layout.renderers.grafico")

_D = {
    "header_bar":  {"left": 0.0, "top": 0.0,  "width": 13.33, "height": 1.2 },
    "title":       {"left": 0.5, "top": 0.15, "width": 12.33, "height": 0.9,  "font_size": 24},
    "accent_line": {"left": 0.0, "top": 1.2,  "width": 13.33, "height": 0.06},
    "chart_area":  {"left": 0.5, "top": 1.4,  "width": 12.33, "height": 5.55},
    "source":      {"left": 0.5, "top": 7.0,  "width": 12.33, "height": 0.4,  "font_size": 10},
}

_CHART_TYPE_MAP: dict[str, XL_CHART_TYPE] = {
    "column": XL_CHART_TYPE.COLUMN_CLUSTERED,
    "bar":    XL_CHART_TYPE.BAR_CLUSTERED,
    "line":   XL_CHART_TYPE.LINE,
    "pie":    XL_CHART_TYPE.PIE,
}

# Series colors — cycles if more than 5 series
_SERIES_COLORS = [
    "1E3A5F",  # primary dark
    "2A9D8F",  # accent teal
    "F4A261",  # orange
    "E76F51",  # coral
    "52B788",  # green
]


def _p(el: str) -> dict[str, Any] | None:
    override = lc.pos("grafico", el)
    if override is None:
        return None
    return {**_D[el], **override}


def _hex_rgb(h: str) -> RGBColor:
    h = h.lstrip("#")
    return RGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))


class GraficoRenderer(BaseLayoutRenderer):
    def render(self, slide: object, model: object, theme: ThemeModel) -> None:
        assert isinstance(model, GraficoSlide)
        c = theme.colors
        f = theme.fonts
        lim = theme.limits

        set_slide_background(slide, c.background)  # type: ignore[arg-type]

        if (p := _p("header_bar")) is not None:
            add_colored_box(slide, p["left"], p["top"], p["width"], p["height"], p.get("fill", c.primary))  # type: ignore[arg-type]

        title = truncate_text(model.title, lim.title_max_chars, "grafico.title")
        if (p := _p("title")) is not None:
            add_text_box(slide, p["left"], p["top"], p["width"], p["height"], title,  # type: ignore[arg-type]
                         f.family, p.get("font_size", f.size_subheading), p.get("color", c.text_light),
                         bold=p.get("bold", f.bold_headings))

        if (p := _p("accent_line")) is not None:
            add_colored_box(slide, p["left"], p["top"], p["width"], p["height"], p.get("fill", c.accent))  # type: ignore[arg-type]

        # Build chart
        chart_data = ChartData()
        chart_data.categories = model.categories
        for series in model.series:
            chart_data.add_series(series.name, tuple(series.values))

        chart_type = _CHART_TYPE_MAP.get(model.chart_type, XL_CHART_TYPE.COLUMN_CLUSTERED)
        if (p := _p("chart_area")) is not None:
            try:
                graphic_frame = slide.shapes.add_chart(  # type: ignore[union-attr]
                    chart_type,
                    Inches(p["left"]), Inches(p["top"]),
                    Inches(p["width"]), Inches(p["height"]),
                    chart_data,
                )
                chart = graphic_frame.chart

                # Apply title if provided
                if model.chart_title:
                    chart.has_title = True
                    chart.chart_title.text_frame.text = model.chart_title
                    for run in chart.chart_title.text_frame.paragraphs[0].runs:
                        run.font.name = f.family
                        run.font.size = Pt(f.size_body + 2)
                        run.font.color.rgb = hex_to_rgb(c.text_dark)
                else:
                    chart.has_title = False

                # Color each series
                if chart_type != XL_CHART_TYPE.PIE:
                    for i, plot_series in enumerate(chart.plots[0].series):
                        color = _SERIES_COLORS[i % len(_SERIES_COLORS)]
                        plot_series.format.fill.solid()
                        plot_series.format.fill.fore_color.rgb = _hex_rgb(color)

                # Clear chart backgrounds (attributes vary by python-pptx version)
                for attr in ("plot_area", "chart_area"):
                    try:
                        getattr(chart, attr).format.fill.background()
                    except Exception:
                        pass

            except Exception as e:
                logger.warning("No se pudo renderizar chart nativo (%s) — omitiendo", e)

        if model.source and (p := _p("source")) is not None:
            add_text_box(slide, p["left"], p["top"], p["width"], p["height"], f"Fuente: {model.source}",  # type: ignore[arg-type]
                         f.family, p.get("font_size", f.size_caption), p.get("color", c.text_muted), italic=True)

        render_layout_extras(slide, "grafico", set(_D.keys()), c.accent, model=model)  # type: ignore[arg-type]
